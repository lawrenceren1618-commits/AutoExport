# -*- coding: utf-8 -*-
"""One-off inspect C26089 order xlsx vs sample if found."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\.cursor\Cursor-projects\AutoExport")
TARGET = Path(
    r"C:\Users\Administrator\AppData\Local\Temp\autoexport-jobs"
    r"\4ed01c57-9446-4dd9-a7fd-488395f6eaed\out\order"
    r"\C07_EF-C26089-C07-40000_order_20260813.xlsx"
)
OUT = ROOT / "2_Processed_AI" / "drafts" / "_c26089_order_inspect.json"


def dump_wb(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    ws = wb.active
    cells = {}
    for row in ws.iter_rows(min_row=1, max_row=50, max_col=12):
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                cells[c.coordinate] = c.value
    merges = [str(m) for m in ws.merged_cells.ranges]
    return {
        "path": str(path),
        "name": path.name,
        "size": path.stat().st_size,
        "sheet": ws.title,
        "cells": cells,
        "merges": merges,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
    }


def find_samples() -> list[str]:
    hits = []
    roots = [
        ROOT / "1_Input_Raw" / "Shipments",
        ROOT / "1_Input_Raw" / "01_Order_Outbound",
        ROOT / "3_Output_Final",
        ROOT / "2_Processed_AI" / "drafts",
        ROOT / "Templates",
    ]
    for base in roots:
        if not base.exists():
            continue
        for p in base.rglob("*"):
            if not p.is_file():
                continue
            n = p.name.lower()
            if p.suffix.lower() not in (".xlsx", ".xls"):
                continue
            if any(k in n for k in ("order", "订单", "c26079", "c26080", "c26081", "备货")):
                hits.append(str(p))
            elif "订单表" in str(p.parent):
                hits.append(str(p))
    return hits[:40]


def main() -> int:
    result: dict = {"target_exists": TARGET.exists(), "samples": find_samples()}
    if TARGET.exists():
        result["target"] = dump_wb(TARGET)
    # also dump draft json
    draft = TARGET.parent / "EF-C26089-C07-40000_order_draft.json"
    if draft.exists():
        result["draft"] = json.loads(draft.read_text(encoding="utf-8"))
    # compare first sample if any
    for s in result["samples"]:
        sp = Path(s)
        # prefer known good EF samples
        if "C26079" in sp.name or "C26080" in sp.name or "订单" in sp.name:
            try:
                result["sample"] = dump_wb(sp)
                result["sample_picked"] = str(sp)
                break
            except Exception as e:
                result.setdefault("sample_errors", []).append(f"{sp}: {e}")
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"wrote {OUT}")
    print(f"target_exists={TARGET.exists()} samples={len(result['samples'])}")
    if result.get("target"):
        print("TARGET CELLS:")
        for k, v in result["target"]["cells"].items():
            print(f"  {k}: {v!r}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
