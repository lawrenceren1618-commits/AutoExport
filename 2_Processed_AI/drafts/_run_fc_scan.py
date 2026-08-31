from __future__ import annotations
import json, re
from datetime import datetime
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(r"D:\.cursor\Cursor-projects\AutoExport")
DRAFTS = ROOT / "2_Processed_AI" / "drafts"

FREIGHT_PATHS = [
    next(ROOT.glob("1_Input_Raw/Shipments/EF-C26040-C04-8490_C04_20260510_8490/货代入库贴下单表/*.xlsx")),
    next(ROOT.glob("1_Input_Raw/Shipments/EF-C26039-C07-10540_C07_20260524_10540/货代入库贴下单表/*C26039*AVP1*.xlsx")),
]
CUSTOMS_PATHS = [
    next(ROOT.glob("1_Input_Raw/Shipments/EF-C26037-C07_C07_20260517_400CTN/报关资料表/*61*报关*.xlsx")),
    next(ROOT.glob("1_Input_Raw/Shipments/EF-C26037-C07_C07_20260517_400CTN/报关资料表/*FBA*.xlsx")),
]

def cell_str(v: Any) -> str:
    if v is None: return ""
    if isinstance(v, float) and v == int(v): return str(int(v))
    if isinstance(v, float): return f"{v:.6g}"
    s = str(v).strip()
    if s.startswith("="): return s
    return s

def dump_sheet(ws, cap_r=85, cap_c=28):
    mr, mc = ws.max_row or 1, ws.max_column or 1
    max_r, max_c = 1, 1
    for row in ws.iter_rows(min_row=1, max_row=mr, min_col=1, max_col=mc):
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                max_r, max_c = max(max_r, cell.row), max(max_c, cell.column)
    max_r, max_c = min(max_r, cap_r), min(max_c, cap_c)
    grid = [{"row": r, "cells": [cell_str(ws.cell(r, c).value) for c in range(1, max_c + 1)]} for r in range(1, max_r + 1)]
    return {"used": {"max_row": max_r, "max_col": max_c}, "grid": grid}

def analyze(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {n: dump_sheet(wb[n]) for n in wb.sheetnames}
    detail_cols = []
    if "下单模板" in wb.sheetnames:
        ws = wb["下单模板"]
        for c in range(1, 25):
            v = ws.cell(14, c).value
            if v: detail_cols.append(str(v).replace("\n", " "))
    wb.close()
    return {
        "path": path.relative_to(ROOT).as_posix(),
        "filename": path.name,
        "sheets": sheets,
        "detail_columns_row14": detail_cols,
    }

def infer_headers(grid):
    best = []
    for gr in grid[:18]:
        cells = [c for c in gr["cells"] if c]
        if len(cells) > len(best):
            best = gr["cells"]
    return [c for c in best if c]

def md_sheet(name, sd):
    lines = [f"### Sheet: `{name}`", "", f"- Used area (capped): rows 1–{sd['used']['max_row']}, cols 1–{sd['used']['max_col']}", ""]
    hdr = infer_headers(sd["grid"])
    if hdr:
        lines += ["**Inferred header / labels:**", ""] + [f"{i}. {c.replace(chr(10),' ')}" for i, c in enumerate(hdr, 1)] + [""]
    ncol = len(sd["grid"][0]["cells"]) if sd["grid"] else 0
    lines += ["**Cell grid:**", "", "| Row | " + " | ".join(f"C{i}" for i in range(1, ncol + 1)) + " |", "|-----|" + "|".join(["---"] * ncol) + "|"]
    for gr in sd["grid"]:
        lines.append("| " + str(gr["row"]) + " | " + " | ".join(c.replace("|", "\\|").replace("\n", "<br>") for c in gr["cells"]) + " |")
    lines.append("")
    return "\n".join(lines)

def notes_freight(a):
    t = json.dumps(a, ensure_ascii=False)
    return {
        "brand_col": "品牌" in t,
        "brands": re.findall(r"无|WIGENIUS|Dressmaker", t)[:8],
        "wh": sorted(set(re.findall(r"\b[A-Z]{3,4}\d\b", a["filename"] + t)))[:12],
        "fba": re.findall(r"FBA[A-Z0-9]+", t)[:6],
    }

FREIGHT_COL_MAP = [
    ("任意订单号（FBA号+U+6位箱号）", "fba_order_carton_ref", "FBA Shipment ID + reference + 箱序（待外箱贴）"),
    ("reference ID", "reference_id", "FBA 外箱贴 / Seller Central"),
    ("箱", "carton_count", "当仓箱数；文件名常带「248箱」"),
    ("中文品名", "product_name_zh", "订单上区"),
    ("英文品名", "product_name_en", "订单上区"),
    ("单箱数量", "qty_per_carton", "订单 条数/箱"),
    ("单价", "declared_unit_price_usd", "业务给定"),
    ("申报币别", "currency", "通常 USD"),
    ("箱重（KG）", "carton_gross_kg", "订单/出货"),
    ("长/宽/高（CM）", "carton_l_w_h_cm", "订单外箱尺寸"),
    ("HS CODE", "hs_code", "6302539090"),
    ("材质", "material", "静态模板文案或固定 polyester"),
    ("用途", "usage", "静态/半静态"),
    ("品牌", "brand", "**固定 `无`**"),
    ("品牌类型", "brand_type", "**固定 `无`**"),
    ("型号", "model", "常 `无`"),
]

fa = [analyze(p) for p in FREIGHT_PATHS]
ca = [analyze(p) for p in CUSTOMS_PATHS]
now = datetime.now().strftime("%Y-%m-%d %H:%M")

fm = [
    "# 货代入库贴下单表 — 已验证样例字段扫描",
    "",
    f"> 生成时间：{now}；检索：`Shipments/**/货代入库贴下单表/`；`1_Input_Raw/02_Freight_Forwarder/` **无 xlsx**；`Templates/` 仅 `order_blank_zh.xlsx`（订单用）。",
    "",
    "## 样例路径（2）",
    "",
]
for a in fa:
    fm.append(f"- `{a['path']}`")
fm += [
    "",
    "## 文档对照（P11 / `02_Freight_Forwarder_Correspondence.md`）",
    "",
    "- 表头 + 明细；品牌列 = **`无`**；HS **6302539090**",
    "- H42：仓库代码 / FBA 地址 / 货件号等待入库货件确认",
    "- 样例结构：**3 sheets** — `下单模板`（主填）、`收货网点`（静态枚举）、`产品名称`（渠道代码枚举）",
    "",
    "## 明细列（row 14，`下单模板` — 两票 C04/C07 一致）",
    "",
]
for i, (cn, key, note) in enumerate(FREIGHT_COL_MAP, 1):
    fm.append(f"{i}. **{cn}** → `{key}` — {note}")
fm.append("")

for a in fa:
    n = notes_freight(a)
    fm += [
        "---",
        f"## `{a['filename']}`",
        f"**路径：** `{a['path']}`",
        f"- **Sheets：** {list(a['sheets'].keys())}",
        f"- **品牌列：** {'有（row14）' if n['brand_col'] else '未检出'}；样例值 {n['brands']}",
        f"- **仓库代码：** {n['wh']}；**FBA 号样例：** {n['fba']}",
        f"- **行数模式：** 表头约 rows 1–13 + **明细 header row 14** + 数据行（样例常 **1 行合并填总箱**；多 SKU 时按注意事项可 **每产品一行** → **row_count 变更**）",
        "",
    ]
    for sn, sd in a["sheets"].items():
        fm.append(md_sheet(sn, sd))
    fm += [
        "**静态 vs 变更（推断）：**",
        "- **不变项：** sheet 版式；`收货网点`/`产品名称` 枚举；列名 row14；HS 默认；品牌/品牌类型/型号填 **`无`**；材质用途模板文案",
        "- **变更项：** 表头区（客户单号、仓库代码、FBA 地址、入库编号、预约链接…）；明细区箱数/单箱数量/件重尺/单价；FBA 号 + reference ID；**row_count**",
        "",
    ]
(DRAFTS / "_field_scan_freight_samples.md").write_text("\n".join(fm), encoding="utf-8")

cm = [
    "# 报关资料表 — 已验证样例字段扫描",
    "",
    f"> 生成时间：{now}；检索：`Shipments/**/报关资料表/`；`03_Customs/`、`3_Output_Final/` **无 xlsx**。",
    "",
    "## 样例路径（2）",
    "",
]
for a in ca:
    cm.append(f"- `{a['path']}`")
cm += [
    "",
    "## 文档对照（P8 / `03_Customs_Correspondence.md`）",
    "",
    "- **5 sheets 套表：** 报关单 / 发票 / 装箱单 / 合同 / 委托协议",
    "- 文件名：`…_{N}箱报关资料` 或含 `YLPSV…=FBA…` 协议号",
    "- C07 样例抬头：**易尚电子商务有限公司**（非订单 C04 页眉口径）",
    "",
]
for a in ca:
    proto = re.findall(r"YLPSV[^=]*|FBA[A-Z0-9]+", a["filename"])
    cm += [
        "---",
        f"## `{a['filename']}`",
        f"**路径：** `{a['path']}`",
        f"- **Sheets：** {list(a['sheets'].keys())}",
        f"- **文件名协议/FBA 片段：** {proto or '（61箱文件无后缀；70箱文件含 YLPSV=FBA）'}",
        f"- **行数模式：** 报关单 ~21 行 + 商品明细区（可变）；发票/装箱单明细 **row_count 随分货 SKU/件数**",
        "",
    ]
    for sn, sd in a["sheets"].items():
        cap = 45 if sn in ("合同", "委托协议") else 85
        if sn in a["sheets"]:
            wb = None
        cm.append(md_sheet(sn, sd))
    cm += [
        "**静态 vs 变更（推断）：**",
        "- **不变项：** 套表制式、海关表头字段标签、卖方/发货人制式段落（模板内）",
        "- **变更项：** 合同协议号、件数/箱数、毛净重、明细数量单价、FBA 标记号、出口日期、**分文件 N 箱**、**row_count**",
        "",
    ]
(DRAFTS / "_field_scan_customs_samples.md").write_text("\n".join(cm), encoding="utf-8")

comb = f"""# 货代 + 报关 — 不变项 / 变更项 对照草案（EF-C26079 / C04-2616）

> 生成时间：{now}；**分析稿**；未改 `settled_knowledge_db.json`；未生成最终 xlsx。

## 当票上下文

| 项 | 值 |
|----|-----|
| 项目夹 | `1_Input_Raw/Shipments/EF-C26079-C04-36460  桌布-2026.7.10/` |
| 订单 | 上 11 行规格 + 下 23 FNSKU；36460 条；945 箱 |
| 出货 | C04-2616；2026.7.31；WIGENIUS / C04 |
| 货代品牌列 | **无** |
| HS | 6302539090 |
| 货代/报关/FBA外箱贴 | **均为空** |

## 扫描样例（verified）

### 货代（2）
"""
for a in fa:
    comb += f"- `{a['path']}`\n"
comb += "\n### 报关（2）\n"
for a in ca:
    comb += f"- `{a['path']}`\n"

comb += """
---

## A. 货代入库贴下单表

### 不变项

| 类别 | 说明 |
|------|------|
| _workbook 结构_ | 3 sheets：`下单模板` + `收货网点` + `产品名称` |
| 列结构 row14 | 见 `_field_scan_freight_samples.md` |
| 品牌 / 品牌类型 / 型号 | 填 **`无`**（样例 C04/C07 均为「无」） |
| HS CODE | `6302539090` |
| 材质/用途模板 | `聚酯纤维/polyester`、`餐桌` 等（可半静态） |
| 渠道枚举 sheets | 不在当票逐格改 |

### 变更项

| 列/区域 | canonical key | 映射来源 |
|---------|---------------|----------|
| 表头：仓库代码 | `warehouse_code` | FBA 外箱贴 / 分仓 |
| 表头：收件地址1 | `dest_address` | 亚马逊仓库地址（样例 HEA2/AVP1 全文） |
| 表头：入库编号 / 预约 | `inbound_ref`, `appointment_*` | Seller Central / 货代 |
| FBA 号 + reference ID | `fba_shipment_id`, `reference_id` | **FBA 外箱贴** |
| 箱（总箱数） | `carton_count` | 分仓箱数（文件名「248箱」「56箱」） |
| 中文/英文品名 | `product_name_zh/en` | 订单上区 |
| 单箱数量 | `qty_per_carton` | 订单 条数/箱 |
| 箱重、长宽高 | `carton_gross_kg`, `carton_dims_cm` | 订单装箱 / 出货 |
| 单价 / 币别 | `declared_unit_price_usd`, `currency` | 业务给定 |
| **row_count** | `freight_detail_rows` | 1 行合并 vs 多 SKU 多行（模板注意事项 #1） |

### 表头静态标签 vs 变更值（`下单模板` rows 2–13）

| 标签（不变） | 样例变更值 |
|--------------|------------|
| 仓库代码 | HEA2 / AVP1 |
| 收货网点 | 金华金东仓 |
| 地址类型 | 亚马逊仓库 |
| 收件公司 | FBA |
| 单证报关 | 是 |
| 收件国家代码 | US |

### 生成判断（C26079）

| 能力 | 现在 | 说明 |
|------|------|------|
| 克隆空白版式 | 部分 | 可用 C04 `EF-C26040…HEA2` 作 **DRAFT 骨架** |
| 写满 H42 / 分仓多文件 | **否** |  blocked：无 FBA 外箱贴、未知仓拆分 |
| 明细多行（11 SKU） | 待确认 | 模板要求多产品「每个产品一列/行」— 与 C04 单行的 248 箱样例不一致 |

---

## B. 报关资料表（5 sheet 套表）

### 不变项

| 类别 | 说明 |
|------|------|
| Sheet 列表 | 报关单 / 发票 / 装箱单 / 合同 / 委托协议 |
| 表头字段名 | 境内发货人、合同协议号、包装种类… |
| HS 默认 | 6302539090（桌布） |

### 变更项

| 列/区域 | canonical key | 映射来源 |
|---------|---------------|----------|
| 合同协议号 | `contract_protocol_no` | P8：`EF-C26037-AVP1` 式；或 `YLPSV…=FBA…` |
| 件数 / 毛净重 | `package_count`, `gross_kg`, `net_kg` | 出货 P7；**每分货文件不同**（61/70/86…） |
| 商品明细 | `customs_lines[]` | 品名规格、数量、单价 |
| 发票 Mark & No | `fba_shipment_id` | 样例 `FBA19CNNVCHX` |
| 出口日期 | `export_date` | 当票日历 |
| 抬头主体 | `declarer_entity` | **须确认 C04 用畅选还是易尚**（样例为 C07 易尚） |
| 退税/征免 | `tax_refund_flag` | **待确认** |
| **row_count** | `customs_line_count` | 发票/装箱单明细行 |

### 生成判断（C26079）

| 能力 | 现在 | 说明 |
|------|------|------|
| 套表 skeleton | 部分 | 仅 C07 历史验证；C04 模板 **待确认** |
| 分货多 xlsx | **否** | blocked：945 箱如何拆分 + 协议号 |
| 申报价 | **否** | 需给定 USD 单价 |

---

## Blockers / 待你确认

1. **FBA 分仓**：945 箱 → 几个仓？是否「一仓一文件（货代+报关）」？
2. **合同协议号**：C26079 用 `EF-C26079-{仓码}` 还是 `YYYYMMDD+序号` / YLPSV 号？
3. **报关模板主体**：C04 当票是否沿用 C07 五件套，还是上传 C04/畅选专用模板到 `03_Customs/`？
4. **报关抬头**：畅选（C04 法律主体）vs 易尚/E-fashion（样例报关单为易尚）？
5. **退税**：36460 条是否全部退税？
6. **货代明细行策略**：945 箱 / 11 SKU — 每 SKU 一行还是每仓合并一行？
7. **Blank 克隆源**：货代默认 `EF-C26040-C04-8490`（C04 单仓）还是等多仓样例？

## FBA 外箱贴依赖

| 单据 | blocked? |
|------|----------|
| 货代下单表（写满） | **是** |
| 报关资料表（分货+协议号） | **是** |
| 仅 DRAFT 空壳 | 可，须标注占位 |

## 参考

- `_field_scan_freight_samples.md` / `_field_scan_customs_samples.md`
- `Templates/02_Freight_Forwarder_Correspondence.md`, `03_Customs_Correspondence.md`, `End_to_End_Shipment_Pipeline.md` (P8/P11)
"""
(DRAFTS / "_change_vs_static_freight_customs_C26079_draft.md").write_text(comb, encoding="utf-8")
print("DONE")