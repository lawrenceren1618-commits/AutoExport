# -*- coding: utf-8 -*-
"""Guard: carton display + C*E*F formulas follow C26079 / packing_standards."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SCRIPTS))

from m1.order_fill import _inner_display, _outer_display, fill_order_template, locate_order_bands
from openpyxl import load_workbook


def test_carton_wording_matches_c26079_row27():
    row = {
        "pack_method": "carton",
        "inner_pack_lwh_cm": "25*19*28",
        "carton_lwh_cm": "40*27*58",
    }
    assert _inner_display(row) == "内盒25*19*28"
    assert _outer_display(row) == "纸箱：40*27*58"


def test_heat_shrink_keeps_spec_cm():
    row = {
        "pack_method": "heat_shrink",
        "inner_pack_lwh_cm": None,
        "carton_lwh_cm": "57*40*33",
    }
    assert _inner_display(row) == "热缩膜包装"
    assert _outer_display(row) == "纸箱规格：57*40*33cm"


def test_fill_writes_formulas_and_carton_text():
    tmp = Path(tempfile.mkdtemp()) / "order.xlsx"
    fill_order_template(
        order_code="EF-C26089-C07-40000",
        order_date="2026-08-10",
        delivery_date="2026-08-29",
        entity_code="C07",
        lines=[
            {
                "size": "60X102inch",
                "color": "白",
                "qty_strips": 40000,
                "qty_cartons": 1000,
                "strips_per_pack": 10,
                "packs_per_carton": 4,
                "fnsku": "X004SJV9UF",
                "gram_weight": "90G",
                "pack_method": "carton",
                "inner_pack_lwh_cm": "25*19*28",
                "carton_lwh_cm": "40*27*58",
                "product_name_zh": "单面磨毛纤维桌布",
            }
        ],
        out_path=tmp,
    )
    ws = load_workbook(tmp).active
    bands = locate_order_bands(ws)
    r = bands["pack_start"]
    assert ws.cell(r, 7).value == f"=C{r}*E{r}*F{r}"
    assert ws.cell(r, 9).value == "内盒25*19*28"
    assert ws.cell(r, 10).value == "纸箱：40*27*58"
    ur = bands["data_start"]
    assert str(ws.cell(ur, 6).value).upper() == f"=G{r}"


if __name__ == "__main__":
    test_carton_wording_matches_c26079_row27()
    test_heat_shrink_keeps_spec_cm()
    test_fill_writes_formulas_and_carton_text()
    print("ok")
