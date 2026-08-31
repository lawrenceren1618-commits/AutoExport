# -*- coding: utf-8 -*-
"""工厂订单：克隆已验证母版，只改变更项；上/下半按变体数伸缩行与合并。"""

from __future__ import annotations

import re
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .canonical import load_entity_rules

PROJECT = Path(__file__).resolve().parent.parent.parent
ORDER_DIR = PROJECT / "1_Input_Raw" / "01_Order_Outbound"
MASTER_XLSX = PROJECT / "Templates" / "order_layout_master.xlsx"
TEMPLATE_XLSX = PROJECT / "Templates" / "order_blank_zh.xlsx"

HEADER_COMPANY = "Xuchang E-fashion Electronic Business Co., Ltd."
HEADER_ADDR = (
    "Add: 7th Floor Building Two, E-commerce Industrial Park Furong Road, "
    "Rural Intergration Demonstration Area of Xuchang, henan, China"
)
DEFAULT_MAKER = "任胜国"
UPPER_COLS = 7
PACK_COLS = 10


def resolve_blank_template() -> Path:
    """退路骨架；正式生成走 resolve_order_template()。"""
    if TEMPLATE_XLSX.exists():
        return TEMPLATE_XLSX
    ensure_blank_template(TEMPLATE_XLSX)
    return TEMPLATE_XLSX


def resolve_order_template() -> Path:
    """权威母版：C26079 订单版式 xlsx。缺文件才退到空白骨架。"""
    if MASTER_XLSX.exists():
        return MASTER_XLSX
    return resolve_blank_template()


def ensure_blank_template(path: Path = TEMPLATE_XLSX) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["D1"] = HEADER_COMPANY
    ws["D2"] = HEADER_ADDR
    ws["D4"] = "TEL:      FAX:    E-MAIL:"
    ws["A6"] = "订单编号："
    ws["D6"] = "下单日期："
    ws["G6"] = "交货日期："
    for col, h in enumerate(
        ["订单编号", "产品品类", "重量", "尺寸规格", "颜色", "数量", "备注"], 1
    ):
        ws.cell(7, col, h)
    ws["A12"] = "制单："
    ws["F12"] = "审核："
    ws["A13"] = "装 箱 单 明 细"
    for col, h in enumerate(
        [
            "尺寸规格",
            "颜色",
            "条数/包",
            "产品标签",
            "包数/箱",
            "总箱数",
            "总条数",
            "箱号",
            "内包装尺寸",
            "外箱尺寸（长*宽*高）",
        ],
        1,
    ):
        ws.cell(14, col, h)
    wb.save(path)
    return path


def format_order_date_cell(iso: str | None) -> str:
    """下单日期单元格：2026.08.10（月日补零，与样例一致）。"""
    if not iso:
        return ""
    parts = re.findall(r"\d+", str(iso))
    if len(parts) < 3:
        return str(iso).replace("-", ".")
    y, m, d = parts[0], parts[1].zfill(2), parts[2].zfill(2)
    if len(y) == 2:
        y = "20" + y
    return f"{y}.{m}.{d}"


def format_delivery_date_cell(iso: str | None) -> str:
    """交货日期：2026.8.21（月日不补零，与样例一致）。"""
    if not iso:
        return ""
    parts = re.findall(r"\d+", str(iso))
    if len(parts) < 3:
        return str(iso).replace("-", ".")
    y, m, d = parts[0], str(int(parts[1])), str(int(parts[2]))
    if len(y) == 2:
        y = "20" + y
    return f"{y}.{m}.{d}"


def order_workbook_filename(order_code: str, delivery_date: str | None) -> str:
    """
    工厂订单文件名（与 Shipments/订单表 样例对齐）：
    `{EF长号}  桌布-{交货日}.xlsx`
    例：EF-C26089-C07-40000  桌布-2026.8.21.xlsx
    禁止前缀店铺码、禁止 `_order_YYYYMMDD` 后缀。
    """
    dd = format_delivery_date_cell(delivery_date) or "待定"
    return f"{order_code}  桌布-{dd}.xlsx"


def size_cell(v: Any) -> str:
    """订单尺寸格：90 圆写母版口径 `90 inch round`。"""
    s = str(v or "").strip()
    compact = re.sub(r"\s+", "", s).lower()
    if "90" in compact and "round" in compact:
        return "90 inch round"
    return s


def gram_display(v: Any) -> str:
    if v is None or v == "":
        return ""
    m = re.search(r"(\d+)", str(v))
    if m:
        return f"{m.group(1)}克/平方米"
    return str(v)


def _is_heat_shrink(row: dict) -> bool:
    if row.get("pack_method") == "heat_shrink":
        return True
    inner = _txt(row.get("inner_pack_lwh_cm"))
    return "热缩" in inner


def _inner_display(row: dict) -> str:
    """沉淀库存 LWH；单元格按 C26079 母版已有写法，不另造「规格/cm」。"""
    if _is_heat_shrink(row):
        return "热缩膜包装"
    if row.get("inner_pack_lwh_cm"):
        ip = str(row["inner_pack_lwh_cm"]).strip()
        return ip if ip.startswith("内盒") else f"内盒{ip}"
    return ""


def _outer_display(row: dict) -> str:
    """
    外箱文字以沉淀库 carton_lwh_cm 为准，套 C26079 已有两类写法：
    - 纸箱/内盒：`纸箱：40*27*58`（无「规格」、不强制 cm）
    - 热缩：`纸箱规格：57*40*33cm`
    禁止把库内纯尺寸一律写成「纸箱规格：…cm」。
    """
    c = row.get("carton_lwh_cm")
    if not c:
        return ""
    s = str(c).strip()
    if s.startswith("纸箱"):
        return s
    if _is_heat_shrink(row):
        core = s[:-2].strip() if s.lower().endswith("cm") else s
        return f"纸箱规格：{core}cm"
    return f"纸箱：{s}"


def _order_short_cell(order_code: str) -> str:
    """明细「订单编号」列：EF-C26089-C07（不含尾条数）。"""
    m = re.match(r"(EF-[A-Z0-9]+-C0[47])", order_code or "", re.I)
    return m.group(1) if m else (order_code or "")


def _txt(v: Any) -> str:
    return str(v or "").strip()


def _compact(v: Any) -> str:
    return re.sub(r"\s+", "", _txt(v))


def _snapshot_pack_styles_by_method(ws, bands: dict) -> dict[str, list[dict]]:
    """样式取自母版同类装箱行（内盒 vs 热缩），禁止一律抄第一行热缩。"""
    start = bands["pack_start"]
    n = max(int(bands.get("pack_n") or 1), 1)
    default = _snapshot_row_styles(ws, start, PACK_COLS)
    out: dict[str, list[dict]] = {"default": default, "carton": None, "heat_shrink": None}
    for r in range(start, start + n):
        inner = _txt(ws.cell(r, 9).value)
        if "热缩" in inner and out["heat_shrink"] is None:
            out["heat_shrink"] = _snapshot_row_styles(ws, r, PACK_COLS)
        if "内盒" in inner and out["carton"] is None:
            out["carton"] = _snapshot_row_styles(ws, r, PACK_COLS)
    out["carton"] = out["carton"] or default
    out["heat_shrink"] = out["heat_shrink"] or default
    return out


def _snapshot_row_styles(ws, row: int, max_col: int) -> list[dict]:
    styles = []
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        styles.append(
            {
                "font": copy(cell.font) if cell.has_style else None,
                "border": copy(cell.border) if cell.has_style else None,
                "fill": copy(cell.fill) if cell.has_style else None,
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment) if cell.has_style else None,
            }
        )
    return styles


def _apply_row_styles(ws, row: int, styles: list[dict]) -> None:
    for c, st in enumerate(styles, 1):
        cell = ws.cell(row, c)
        if st.get("font") is not None:
            cell.font = st["font"]
        if st.get("border") is not None:
            cell.border = st["border"]
        if st.get("fill") is not None:
            cell.fill = st["fill"]
        if st.get("number_format"):
            cell.number_format = st["number_format"]
        if st.get("alignment") is not None:
            cell.alignment = st["alignment"]


def _set_value_keep_style(cell, value) -> None:
    """改文字，保留原字体/填充/对齐（页眉日期、订单号）。"""
    font = copy(cell.font) if cell.has_style else None
    fill = copy(cell.fill) if cell.has_style else None
    border = copy(cell.border) if cell.has_style else None
    al = copy(cell.alignment) if cell.has_style else None
    nf = cell.number_format
    cell.value = value
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if al is not None:
        cell.alignment = al
    if nf:
        cell.number_format = nf


def _unmerge_from(ws, start_row: int) -> None:
    to_unmerge = []
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= start_row:
            to_unmerge.append(str(rng))
    for rng in to_unmerge:
        try:
            ws.unmerge_cells(rng)
        except ValueError:
            pass


def _merge(ws, r1: int, c1: int, r2: int, c2: int) -> None:
    if r2 < r1 or c2 < c1:
        return
    if r1 == r2 and c1 == c2:
        return
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(r1, c1)
    cell.alignment = Alignment(
        horizontal=cell.alignment.horizontal or "center",
        vertical="center",
        wrap_text=True if cell.alignment.wrap_text else cell.alignment.wrap_text,
    )


def _center(ws, row: int, col: int) -> None:
    cell = ws.cell(row, col)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


def locate_order_bands(ws) -> dict[str, int]:
    """按标签定位上半/下半，不依赖母版固定行数。"""
    upper_header = None
    for r in range(1, min(ws.max_row or 1, 30) + 1):
        a, b = _txt(ws.cell(r, 1).value), _txt(ws.cell(r, 2).value)
        if a == "订单编号" and "品类" in b:
            upper_header = r
            break
    if not upper_header:
        raise ValueError("order master missing upper header 订单编号/产品品类")

    data_start = upper_header + 1
    upper_total = None
    for r in range(data_start, (ws.max_row or data_start) + 1):
        if "合计" in _compact(ws.cell(r, 1).value):
            upper_total = r
            break
    if not upper_total:
        raise ValueError("order master missing upper 合计")

    maker = None
    pack_title = None
    pack_header = None
    for r in range(upper_total + 1, (ws.max_row or upper_total) + 1):
        a = _txt(ws.cell(r, 1).value)
        ac = _compact(a)
        if maker is None and ac.startswith("制单"):
            maker = r
        if "装箱" in ac:
            pack_title = r
        c3 = _compact(ws.cell(r, 3).value)
        if ac == "尺寸规格" and ("条数" in c3 or "包" in c3):
            pack_header = r
            break
    if not pack_header:
        raise ValueError("order master missing packing header")

    pack_start = pack_header + 1
    pack_total = None
    for r in range(pack_start, (ws.max_row or pack_start) + 1):
        if "合计" in _compact(ws.cell(r, 1).value):
            pack_total = r
            break
    if not pack_total:
        raise ValueError("order master missing packing 合计")

    return {
        "upper_header": upper_header,
        "data_start": data_start,
        "upper_total": upper_total,
        "upper_n": upper_total - data_start,
        "maker": maker or (upper_total + 2),
        "pack_title": pack_title or (pack_header - 1),
        "pack_header": pack_header,
        "pack_start": pack_start,
        "pack_total": pack_total,
        "pack_n": pack_total - pack_start,
    }


def resize_band(
    ws,
    *,
    data_start: int,
    old_n: int,
    new_n: int,
    styles: list[dict],
    max_col: int,
    row_height: float | None,
) -> None:
    """伸缩数据行：insert/delete 发生在合计行之前，下方制单/装箱区随表格移动。"""
    if new_n < 1:
        raise ValueError("data band needs at least 1 row")
    if old_n < 1:
        old_n = 1
    if new_n > old_n:
        insert_at = data_start + old_n
        ws.insert_rows(insert_at, amount=new_n - old_n)
        for i in range(new_n - old_n):
            r = insert_at + i
            _apply_row_styles(ws, r, styles)
            if row_height:
                ws.row_dimensions[r].height = row_height
    elif new_n < old_n:
        ws.delete_rows(data_start + new_n, old_n - new_n)


def _clear_band(ws, start: int, n: int, max_col: int) -> None:
    for r in range(start, start + n):
        for c in range(1, max_col + 1):
            ws.cell(r, c).value = None


def _summarize_upper(lines: list[dict]) -> list[dict]:
    summary: dict[tuple, dict] = {}
    for ln in lines:
        k = (ln.get("size"), ln.get("color"))
        if k not in summary:
            summary[k] = {
                "size": ln.get("size"),
                "color": ln.get("color"),
                "gram_weight": ln.get("gram_weight"),
                "product_name_zh": ln.get("product_name_zh") or "单面磨毛纤维桌布",
                "qty_strips": 0.0,
            }
        summary[k]["qty_strips"] += float(ln.get("qty_strips") or 0)
    out = []
    for s in summary.values():
        qty = s["qty_strips"]
        s["qty_strips"] = int(qty) if qty == int(qty) else qty
        out.append(s)
    return out


def _size_merge_spans(rows: list[dict], start: int) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    i = 0
    while i < len(rows):
        j = i + 1
        while j < len(rows) and rows[j].get("size") == rows[i].get("size"):
            j += 1
        spans.append((start + i, start + j - 1))
        i = j
    return spans


def _thin_border_if_empty(ws, r1: int, r2: int, max_col: int) -> None:
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for r in range(r1, r2 + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if cell.border is None or cell.border.left is None or cell.border.left.style is None:
                cell.border = thin


def fill_order_template(
    *,
    order_code: str,
    order_date: str | None,
    delivery_date: str | None,
    entity_code: str,
    lines: list[dict],
    out_path: Path | str,
    maker: str | None = None,
) -> Path:
    """
    克隆 Templates/order_layout_master.xlsx（C26079 已验证版式）：
    只改字段与行数；上半=尺寸+颜色一种一行；下半=包规/FNSKU 变体一行。
    行数可多于或少于母版，insert/delete 后重做合并，不套死母版行数。
    下半总条数写 =C*E*F；上半数量引用对应 G。内外箱按沉淀库 LWH + 母版已有文案。
    """
    rules = load_entity_rules()
    ent = rules.get("entities", {}).get(entity_code, {})
    _ = ent.get("brand_tablecloth")

    if not lines:
        raise ValueError("no order lines")

    upper = _summarize_upper(lines)
    pack_lines = list(lines)
    n_u, n_p = len(upper), len(pack_lines)

    tpl = resolve_order_template()
    wb = load_workbook(tpl)
    ws = wb.active

    # 页眉只改可变文字；公司名/字体保持母版（Tahoma），不改成粗体 14
    _set_value_keep_style(ws["A6"], f"订单编号：{order_code}")
    _set_value_keep_style(ws["E6"], format_order_date_cell(order_date))
    _set_value_keep_style(ws["G6"], f"交货日期：{format_delivery_date_cell(delivery_date)}")

    bands = locate_order_bands(ws)
    upper_styles = _snapshot_row_styles(ws, bands["data_start"], UPPER_COLS)
    pack_styles_by = _snapshot_pack_styles_by_method(ws, bands)
    pack_styles = pack_styles_by["default"]
    upper_header_styles = _snapshot_row_styles(ws, bands["upper_header"], UPPER_COLS)
    upper_total_styles = _snapshot_row_styles(ws, bands["upper_total"], UPPER_COLS)
    pack_header_styles = _snapshot_row_styles(ws, bands["pack_header"], PACK_COLS)
    pack_total_styles = _snapshot_row_styles(ws, bands["pack_total"], PACK_COLS)
    uh = ws.row_dimensions[bands["data_start"]].height
    ph = ws.row_dimensions[bands["pack_start"]].height
    uth = ws.row_dimensions[bands["upper_total"]].height
    pth = ws.row_dimensions[bands["pack_total"]].height
    uhh = ws.row_dimensions[bands["upper_header"]].height
    phh = ws.row_dimensions[bands["pack_header"]].height

    _unmerge_from(ws, bands["data_start"])

    resize_band(
        ws,
        data_start=bands["data_start"],
        old_n=bands["upper_n"],
        new_n=n_u,
        styles=upper_styles,
        max_col=UPPER_COLS,
        row_height=uh,
    )
    bands = locate_order_bands(ws)
    resize_band(
        ws,
        data_start=bands["pack_start"],
        old_n=bands["pack_n"],
        new_n=n_p,
        styles=pack_styles,
        max_col=PACK_COLS,
        row_height=ph,
    )
    bands = locate_order_bands(ws)

    _clear_band(ws, bands["data_start"], n_u, UPPER_COLS)
    _clear_band(ws, bands["pack_start"], n_p, PACK_COLS)

    short = _order_short_cell(order_code)
    first = upper[0]
    total_strips = 0.0
    for i, s in enumerate(upper):
        r = bands["data_start"] + i
        _apply_row_styles(ws, r, upper_styles)
        if uh:
            ws.row_dimensions[r].height = uh
        qty = s["qty_strips"]
        total_strips += float(qty or 0)
        if i == 0:
            ws.cell(r, 1, short)
            ws.cell(r, 2, first["product_name_zh"])
            ws.cell(r, 3, gram_display(first.get("gram_weight")))
        ws.cell(r, 4, size_cell(s.get("size")))
        ws.cell(r, 5, s.get("color"))
        # 上半数量稍后写成 =Gxx（对齐 C26079 F9=G27），此处先占位

    ur0, ur1 = bands["data_start"], bands["data_start"] + n_u - 1
    tot_u = bands["upper_total"]
    _apply_row_styles(ws, bands["upper_header"], upper_header_styles)
    if uhh:
        ws.row_dimensions[bands["upper_header"]].height = uhh
    _apply_row_styles(ws, tot_u, upper_total_styles)
    if uth:
        ws.row_dimensions[tot_u].height = uth
    ws.cell(tot_u, 1, "合   计")
    for c in range(2, UPPER_COLS + 1):
        if c != 6:
            ws.cell(tot_u, c, None)
    ws.cell(tot_u, 6, f"=SUM({get_column_letter(6)}{ur0}:{get_column_letter(6)}{ur1})")

    maker_row = bands["maker"]
    ws.cell(maker_row, 1, f"制单：{maker or DEFAULT_MAKER}")
    ws.cell(maker_row, 6, "审核：")

    pack_cartons = 0.0
    pack_strips = 0.0
    g_by_kind: dict[tuple, list[str]] = {}
    for i, ln in enumerate(pack_lines):
        r = bands["pack_start"] + i
        method_key = "heat_shrink" if _is_heat_shrink(ln) else "carton"
        _apply_row_styles(ws, r, pack_styles_by.get(method_key) or pack_styles)
        if ph:
            ws.row_dimensions[r].height = ph
        packs = ln.get("packs_per_carton")
        if packs in (None, "") and ln.get("strips_per_carton") and ln.get("strips_per_pack"):
            spp = int(float(ln["strips_per_pack"]))
            spc = int(float(ln["strips_per_carton"]))
            if spp > 0 and spc % spp == 0:
                packs = spc // spp
        ws.cell(r, 1, size_cell(ln.get("size") or ""))
        ws.cell(r, 2, ln.get("color") or "")
        ws.cell(r, 3, ln.get("strips_per_pack") or "")
        ws.cell(r, 4, ln.get("fnsku") or "")
        ws.cell(r, 5, packs if packs is not None else "")
        ws.cell(r, 6, ln.get("qty_cartons") or "")
        # C26079：总条数 = 条数/包 × 包数/箱 × 总箱数
        ws.cell(r, 7, f"=C{r}*E{r}*F{r}")
        ws.cell(r, 8, ln.get("carton_range") or "")
        ws.cell(r, 9, _inner_display(ln))
        ws.cell(r, 10, _outer_display(ln))
        pack_cartons += float(ln.get("qty_cartons") or 0)
        pack_strips += float(ln.get("qty_strips") or 0)
        g_by_kind.setdefault((ln.get("size"), ln.get("color")), []).append(f"G{r}")

    for i, s in enumerate(upper):
        r = bands["data_start"] + i
        refs = g_by_kind.get((s.get("size"), s.get("color"))) or []
        if refs:
            ws.cell(r, 6, "=" + "+".join(refs))
        else:
            ws.cell(r, 6, s.get("qty_strips") or "")

    pr0, pr1 = bands["pack_start"], bands["pack_start"] + n_p - 1
    tot_p = bands["pack_total"]
    _apply_row_styles(ws, bands["pack_header"], pack_header_styles)
    if phh:
        ws.row_dimensions[bands["pack_header"]].height = phh
    _apply_row_styles(ws, tot_p, pack_total_styles)
    if pth:
        ws.row_dimensions[tot_p].height = pth
    ws.cell(tot_p, 1, "合计")
    for c in range(2, PACK_COLS + 1):
        if c not in (6, 7):
            ws.cell(tot_p, c, None)
    ws.cell(tot_p, 6, f"=SUM(F{pr0}:F{pr1})")
    ws.cell(tot_p, 7, f"=SUM(G{pr0}:G{pr1})")

    if n_u > 1:
        _merge(ws, ur0, 1, ur1, 1)
        _merge(ws, ur0, 2, ur1, 2)
        _merge(ws, ur0, 3, ur1, 3)
    _center(ws, ur0, 1)
    _center(ws, ur0, 2)
    _center(ws, ur0, 3)
    for a, b in _size_merge_spans(upper, ur0):
        _merge(ws, a, 4, b, 4)
        _center(ws, a, 4)

    _merge(ws, tot_u, 1, tot_u, 5)
    _merge(ws, bands["pack_title"], 1, bands["pack_title"], 9)
    _merge(ws, tot_p, 1, tot_p, 5)

    _thin_border_if_empty(ws, ur0, tot_u, UPPER_COLS)
    _thin_border_if_empty(ws, bands["pack_header"], tot_p, PACK_COLS)

    last = max(tot_p, ws.max_row or tot_p)
    ws.print_area = f"A1:{get_column_letter(PACK_COLS)}{last}"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
