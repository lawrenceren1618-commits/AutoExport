# -*- coding: utf-8 -*-
"""出货表填空（按 C04-2615 样本布局；无 FNSKU 列）。

版式强制：
- 数据区自动扩行，行高统一
- 产品品类 / 重量 / 所属订单 / 目的地 纵向合并
- 所属订单 = 订单短码 Cxxxxx（样例 C26064）；所属仓库仍空
- 合计用 =SUM(E…)
- 合计行以下无框线（清掉模板残留空表格边框）
"""

from __future__ import annotations

import re
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parent.parent.parent
TEMPLATE_CANDIDATES = [
    PROJECT
    / "1_Input_Raw"
    / "Shipments"
    / "C04-2615-U_C04_20260624_TBD"
    / "出货表"
    / "C04-2615-U-2026.6.24.xlsx",
    PROJECT / "Templates" / "outbound_blank_zh.xlsx",
]

DATA_START = 5
DEFAULT_ROW_HEIGHT = 37.0


def resolve_outbound_template() -> Path:
    for p in TEMPLATE_CANDIDATES:
        if p.is_file():
            return p
    raise FileNotFoundError(
        "outbound template missing; expected C04-2615 sample under Shipments/出货表"
    )


def format_ship_date_cell(iso: str | None) -> str:
    """2026-07-31 → 2026.7.31"""
    if not iso:
        return ""
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(iso).strip())
    if not m:
        return str(iso)
    return f"{m.group(1)}.{int(m.group(2))}.{int(m.group(3))}"


def format_outbound_filename(short: str, ship_date: str | None) -> str:
    return f"{short}-U-{format_ship_date_cell(ship_date)}.xlsx"


def order_short_for_belonging(order_code: str | None, order_short: str | None = None) -> str:
    """所属订单：优先显式短码；否则从 EF-C26080-… 抽 C26080。"""
    if order_short:
        s = str(order_short).strip()
        if re.fullmatch(r"C\d{5}", s, flags=re.I):
            return s.upper()
        m = re.search(r"(C\d{5})", s, flags=re.I)
        if m:
            return m.group(1).upper()
    if order_code:
        m = re.search(r"EF-(C\d{5})", str(order_code), flags=re.I)
        if m:
            return m.group(1).upper()
        m = re.search(r"\b(C\d{5})\b", str(order_code), flags=re.I)
        if m:
            return m.group(1).upper()
    return ""


def summarize_size_color(lines: list[dict]) -> list[dict]:
    summary: dict[tuple, dict] = {}
    for ln in lines:
        size = ln.get("size")
        color = ln.get("color")
        if not size and not color:
            continue
        k = (size, color)
        if k not in summary:
            summary[k] = {
                "size": size,
                "color": color,
                "gram_weight": ln.get("gram_weight") or "90G",
                "product_name_zh": ln.get("product_name_zh") or "单面磨毛纤维桌布",
                "qty_strips": 0.0,
            }
        summary[k]["qty_strips"] += float(ln.get("qty_strips") or 0)
    rows = []
    for (_sz, _cl), s in summary.items():
        qty = s["qty_strips"]
        rows.append(
            {
                **s,
                "qty_strips": int(qty) if qty == int(qty) else qty,
            }
        )
    return rows


def _copy_cell_style(src, dst) -> None:
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def _snapshot_row_styles(ws, row: int, max_col: int = 9) -> list[Any]:
    styles = []
    for c in range(1, max_col + 1):
        cell = ws.cell(row, c)
        styles.append(
            {
                "font": copy(cell.font) if cell.has_style else None,
                "border": copy(cell.border) if cell.has_style else None,
                "fill": copy(cell.fill) if cell.has_style else None,
                "number_format": cell.number_format,
                "alignment": copy(cell.alignment) if cell.has_style else None,
            }
        )
    return styles


def _apply_row_styles(ws, row: int, styles: list[Any]) -> None:
    for c, st in enumerate(styles, 1):
        cell = ws.cell(row, c)
        if st.get("font") is not None:
            cell.font = st["font"]
        if st.get("border") is not None:
            cell.border = st["border"]
        if st.get("fill") is not None:
            cell.fill = st["fill"]
        if st.get("number_format"):
            cell.number_format = st["number_format"]
        if st.get("alignment") is not None:
            cell.alignment = st["alignment"]


def _unmerge_from(ws, start_row: int) -> None:
    to_unmerge = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= start_row or (rng.min_row <= start_row <= rng.max_row):
            # unmerge any range touching data band+
            if rng.max_row >= start_row:
                to_unmerge.append(str(rng))
    for rng in to_unmerge:
        try:
            ws.unmerge_cells(rng)
        except ValueError:
            pass


_NO_BORDER = Border(
    left=Side(style=None),
    right=Side(style=None),
    top=Side(style=None),
    bottom=Side(style=None),
)


def _clear_rows(ws, start_row: int, end_row: int, max_c: int = 9) -> None:
    for r in range(start_row, end_row + 1):
        for c in range(1, max_c + 1):
            ws.cell(r, c).value = None


def _clear_borders(ws, start_row: int, end_row: int, max_c: int = 9) -> None:
    """合计行以下（或整段空行）去掉框线，避免模板残留空表格。"""
    if end_row < start_row:
        return
    for r in range(start_row, end_row + 1):
        for c in range(1, max_c + 1):
            ws.cell(r, c).border = copy(_NO_BORDER)


def _gram_display(gw: Any) -> str:
    if isinstance(gw, (int, float)):
        return f"{int(gw)}克/平方米" if float(gw) == int(gw) else f"{gw}克/平方米"
    gs = str(gw or "90G")
    if "克" in gs:
        return gs
    m = re.match(r"^(\d+)\s*G$", gs, flags=re.I)
    if m:
        return f"{m.group(1)}克/平方米"
    return gs if gs else "90克/平方米"


def fill_outbound_template(
    *,
    outbound_short: str,
    ship_date: str | None,
    lines: list[dict],
    out_path: Path | str,
    sales_person: str = "任胜国",
    destination: str = "US",
    template_path: Path | None = None,
    order_code: str | None = None,
    order_short: str | None = None,
) -> Path:
    """
    填出货表：页眉短号+发货日；明细=尺寸+颜色汇总；
    所属订单=订单短码；所属仓库留空；品类/克重/所属订单/目的地合并。
    """
    detail = summarize_size_color(lines)
    if not detail:
        raise ValueError("no outbound detail lines (need size/color/qty_strips)")

    n = len(detail)
    total = sum(float(r["qty_strips"] or 0) for r in detail)
    total_i = int(total) if total == int(total) else total
    belong = order_short_for_belonging(order_code, order_short)

    tpl = template_path or resolve_outbound_template()
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(tpl)
    ws = wb.active

    # Header
    ws["A1"] = "发货单"
    ws["A2"] = "订单发货日期："
    ws["B2"] = format_ship_date_cell(ship_date)
    ws["F2"] = "客户及订单号："
    ws["G2"] = outbound_short
    ws["A3"] = "订单数量："
    ws["B3"] = total_i
    ws["F3"] = "业务员："
    ws["G3"] = sales_person

    headers = [
        "产品品类",
        "重量",
        "尺寸规格",
        "颜色",
        "数量",
        "所属订单",
        "所属仓库",
        "备注",
        "目的地",
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(4, col, h)

    # Snapshot style from template data row before clearing
    style_row = DATA_START
    row_styles = _snapshot_row_styles(ws, style_row, 9)
    title_merge = None
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == 1 and rng.max_row == 1:
            title_merge = str(rng)
            break

    _unmerge_from(ws, DATA_START)
    # Clear a generous band (template + extra)
    old_max = max(ws.max_row or DATA_START, DATA_START + 20)
    _clear_rows(ws, DATA_START, old_max, 9)

    data_end = DATA_START + n - 1
    total_row = data_end + 1

    # Ensure Excel physically has enough rows: insert if needed
    # (openpyxl writes beyond max_row fine; insert keeps print area habits)
    existing_band = 9  # sample had 9 detail rows
    if n > existing_band:
        ws.insert_rows(DATA_START + existing_band, amount=n - existing_band)

    product = detail[0].get("product_name_zh") or "单面磨毛纤维桌布"
    gw_disp = _gram_display(detail[0].get("gram_weight"))

    for i, row in enumerate(detail):
        r = DATA_START + i
        _apply_row_styles(ws, r, row_styles)
        ws.row_dimensions[r].height = DEFAULT_ROW_HEIGHT
        if i == 0:
            ws.cell(r, 1, product)
            ws.cell(r, 2, gw_disp)
            ws.cell(r, 6, belong or None)
            ws.cell(r, 9, destination)
        else:
            ws.cell(r, 1, None)
            ws.cell(r, 2, None)
            ws.cell(r, 6, None)
            ws.cell(r, 9, None)
        ws.cell(r, 3, row.get("size"))
        ws.cell(r, 4, row.get("color"))
        ws.cell(r, 5, row.get("qty_strips"))
        ws.cell(r, 7, None)  # 所属仓库
        ws.cell(r, 8, None)

    # Total row
    _apply_row_styles(ws, total_row, row_styles)
    ws.row_dimensions[total_row].height = DEFAULT_ROW_HEIGHT
    ws.cell(total_row, 1, "合计")
    for c in range(2, 10):
        if c != 5:
            ws.cell(total_row, c, None)
    ws.cell(total_row, 5, f"=SUM(E{DATA_START}:E{data_end})")

    # Vertical merges: 品类 / 重量 / 所属订单 / 目的地
    if n >= 1:
        merge_cols = [1, 2, 6, 9]
        for c in merge_cols:
            if n > 1:
                ws.merge_cells(
                    start_row=DATA_START,
                    start_column=c,
                    end_row=data_end,
                    end_column=c,
                )
            cell = ws.cell(DATA_START, c)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # Keep title merge if present
    if title_merge and title_merge not in {str(x) for x in ws.merged_cells.ranges}:
        try:
            ws.merge_cells(title_merge)
        except ValueError:
            pass

    # Thin border fallback on data band if template style missing
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for r in range(DATA_START, total_row + 1):
        for c in range(1, 10):
            cell = ws.cell(r, c)
            if cell.border is None or cell.border.left.style is None:
                cell.border = thin

    # 合计行以下不得留框线（模板多明细时空行会残留边框）
    below = max(old_max, ws.max_row or total_row, total_row + 20)
    _clear_borders(ws, total_row + 1, below, 9)

    wb.save(out)
    return out


def refill_outbound_file(
    src_path: Path | str,
    *,
    order_code: str | None = None,
    order_short: str | None = None,
    out_path: Path | str | None = None,
) -> Path:
    """从已有出货表读明细行，按正确版式重写（用于修复历史坏文件）。"""
    src = Path(src_path)
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.active
    outbound_short = str(ws["G2"].value or "").strip()
    ship_raw = str(ws["B2"].value or "").strip()
    # 2026.7.31 → 2026-07-31
    ship_date = None
    m = re.match(r"^(\d{4})\.(\d{1,2})\.(\d{1,2})$", ship_raw)
    if m:
        ship_date = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"

    lines: list[dict] = []
    # find header row
    header_row = 4
    r = DATA_START
    while r <= (ws.max_row or DATA_START):
        a = ws.cell(r, 1).value
        if a == "合计":
            break
        size = ws.cell(r, 3).value
        color = ws.cell(r, 4).value
        qty = ws.cell(r, 5).value
        if size or color or qty:
            # product/weight from first non-empty in band
            product = None
            gw = None
            for rr in range(DATA_START, r + 1):
                if ws.cell(rr, 1).value and ws.cell(rr, 1).value != "合计":
                    product = ws.cell(rr, 1).value
                    break
            for rr in range(DATA_START, r + 1):
                if ws.cell(rr, 2).value:
                    gw = ws.cell(rr, 2).value
                    break
            lines.append(
                {
                    "size": size,
                    "color": color,
                    "qty_strips": qty,
                    "product_name_zh": product,
                    "gram_weight": gw,
                }
            )
        r += 1

    if not lines:
        raise ValueError(f"no detail lines in {src}")

    dest = out_path or src
    return fill_outbound_template(
        outbound_short=outbound_short or "C00-0000",
        ship_date=ship_date,
        lines=lines,
        out_path=dest,
        order_code=order_code,
        order_short=order_short,
    )
