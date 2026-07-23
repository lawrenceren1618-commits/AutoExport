# -*- coding: utf-8 -*-
"""Fix C26079-B AWD-IUTE freight xlsx: restore DISPIMG + xl/media from template."""
from __future__ import annotations

import shutil
import sys
import zipfile
import re
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:/.cursor/Cursor-projects/AutoExport")
sys.path.insert(0, str(ROOT / "Scripts"))
import fix_c26079_a_freight_xlsx as fixa

TEMPLATE = fixa.TEMPLATE
DISPIMG = fixa.DISPIMG_FORMULA
C04_URL = fixa.C04_SALES_URL
IUTE_ADDR = "IUTE - AWD 3PL - 16155 W Elwood St - GOODYEAR, AZ 85338 - United States"

FILENAME = "货代下单表_C26079-B 2026.7.31桌布 AWD-IUTE 140箱.xlsx"


def freight_path() -> Path:
    d = fixa.find_freight_dir()
    return d / FILENAME


def inspect(path: Path) -> tuple[int, str]:
    n, media = fixa.count_media(path)
    s15_snip = "(no file)"
    if path.is_file():
        with zipfile.ZipFile(path, "r") as z:
            for sf in z.namelist():
                if "sheet" in sf and sf.endswith(".xml"):
                    data = z.read(sf).decode("utf-8", errors="replace")
                    if 'r="S15"' in data:
                        m = re.search(r"<c r=\"S15\"[^>]*>(.*?)</c>", data, re.DOTALL)
                        s15_snip = (m.group(0)[:240] if m else "?")
                        break
    return n, s15_snip


def apply_values(ws) -> None:
    ws["C2"] = "111111111"
    ws["E3"] = "IUTE"
    ws["E7"] = None
    ws["E8"] = IUTE_ADDR
    for r in ("E10", "E11", "E12"):
        ws[r] = None
    ws["A15"] = "STAR-QG6CTURZGY6MX"
    ws["B15"] = None
    ws["C15"] = 140
    ws["F15"] = 30
    ws["G15"] = 3.63
    fixa.apply_row15_fixed(ws)
    s15 = ws["S15"].value
    if not (isinstance(s15, str) and "DISPIMG" in s15.upper()):
        ws["S15"] = DISPIMG
    ws["T15"] = C04_URL


def fix(dest: Path | None = None) -> int:
    dest = dest or freight_path()
    if not TEMPLATE.is_file():
        print(f"ERROR: template missing: {TEMPLATE}")
        return 1
    before_n, before_s15 = inspect(dest)
    print(f"BEFORE media={before_n}")
    print(f"BEFORE S15: {before_s15}")

    tmp = dest.with_name("_C26079-B_mediafix_tmp.xlsx")
    shutil.copy2(TEMPLATE, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb.active
    apply_values(ws)
    wb.save(tmp)
    wb.close()
    reinjected = fixa.reinject_media_from_template(tmp, TEMPLATE)

    after_n, after_s15 = inspect(tmp)
    print(f"TMP AFTER media={after_n} reinjected={reinjected}")
    print(f"TMP AFTER S15: {after_s15}")

    try:
        tmp.replace(dest)
        print(f"Wrote: {dest}")
    except OSError as e:
        print(f"WARN: could not replace locked target ({e})")
        print(f"Fixed copy ready at: {tmp}")
        print("Close Excel on the target file and re-run this script.")
        return 2

    final_n, final_s15 = inspect(dest)
    print(f"AFTER media={final_n}")
    print(f"AFTER S15: {final_s15}")
    return 0 if final_n >= 1 else 3


if __name__ == "__main__":
    raise SystemExit(fix())
