# -*- coding: utf-8 -*-
"""One-off fix for C26079-A freight forwarder xlsx (5 warehouses)."""
from __future__ import annotations

import math
import re
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"D:/.cursor/Cursor-projects/AutoExport")
TEMPLATE = (
    ROOT
    / "1_Input_Raw/Shipments/EF-C26040-C04-8490_C04_20260510_8490/货代入库贴下单表/货代入库贴数据2026.5.10 HEA2 248箱.xlsx"
)
NOTE_PATH = ROOT / "2_Processed_AI/drafts/C26079-A_freight_fix_note.md"

C04_SALES_URL = "https://www.amazon.com/dp/B0CCRKDW1K?th=1"

DISPIMG_FORMULA = '=DISPIMG("ID_3D20B9B273A54ACBA2D957C7EE1766FF",1)'

WAREHOUSES = {
    "AVP1": {
        "fba": "FBA19JT39521",
        "track": "36AGPMRC",
        "cartons": 175,
        "strips": 6966,
        "unit": 3.18,
        "addr": "AVP1 - 550 Oak Ridge Road - Hazle Township, PA 18202-9361 - United States",
    },
    "CLT2": {
        "fba": "FBA19JT71WTZ",
        "track": "629YFDRT",
        "cartons": 128,
        "strips": 5136,
        "unit": 3.16,
        "addr": "CLT2 - 10240 Old Dowd Rd - Charlotte, NC 28214-8082 - United States",
    },
    "GYR2": {
        "fba": "FBA19JT3Q0H0",
        "track": "2C4TSS5H",
        "cartons": 139,
        "strips": 5608,
        "unit": 3.15,
        "addr": "GYR2 - 17341 W MINNEZONA AVE - GOODYEAR, AZ 85395-7672 - United States",
    },
    "LGB8": {
        "fba": "FBA19JT0K3PY",
        "track": "6522FTZR",
        "cartons": 182,
        "strips": 7366,
        "unit": 3.12,
        "addr": "LGB8 - 1568 N Linden Ave - Rialto, CA 92376-8624 - United States",
    },
    "MDW2": {
        "fba": "FBA19JT16DMQ",
        "track": "25HS4ZWF",
        "cartons": 181,
        "strips": 7184,
        "unit": 3.17,
        "addr": "MDW2 - 250 EMERALD DR - Joliet, IL 60433-3280 - United States",
    },
}

KEY_CELLS = [
    "C2",
    "C3",
    "C5",
    "E3",
    "E7",
    "E8",
    "E10",
    "E11",
    "E12",
    "A15",
    "B15",
    "C15",
    "D15",
    "E15",
    "F15",
    "G15",
    "H15",
    "I15",
    "J15",
    "K15",
    "L15",
    "M15",
    "N15",
    "O15",
    "P15",
    "Q15",
    "R15",
    "S15",
    "T15",
]

MEDIA_PARTS = (
    "xl/cellimages.xml",
    "xl/_rels/cellimages.xml.rels",
)


def find_freight_dir() -> Path:
    shipments = ROOT / "1_Input_Raw/Shipments"
    for project in shipments.iterdir():
        if not project.is_dir() or "C26079" not in project.name:
            continue
        for sub in project.iterdir():
            if not sub.is_dir():
                continue
            for xlsx in sub.glob("*.xlsx"):
                if xlsx.name.startswith("~$"):
                    continue
                if "C26079-A" in xlsx.name and xlsx.name.startswith("货代下单表"):
                    return sub
    raise FileNotFoundError("C26079-A freight directory not found under Shipments")


def read_template_c2() -> str:
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)
    try:
        val = wb.active["C2"].value
    finally:
        wb.close()
    if val is None:
        raise ValueError("Template C2 is empty; expected placeholder customer number")
    return str(val).strip()


def warehouse_from_name(name: str) -> str:
    for code in WAREHOUSES:
        if code in name:
            return code
    raise ValueError(f"Cannot infer warehouse code from filename: {name}")


def cell_str(value) -> str:
    if value is None:
        return "(empty)"
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value)


def count_media(path: Path) -> tuple[int, list[str]]:
    with zipfile.ZipFile(path, "r") as zf:
        media = [n for n in zf.namelist() if n.startswith("xl/media/") and not n.endswith("/")]
    return len(media), media


def needs_media_reinject(path: Path) -> bool:
    n, _ = count_media(path)
    return n == 0


def merge_xml_snippet(xml: str, snippet: str) -> str:
    if snippet in xml:
        return xml
    return xml.replace("</Types>", snippet + "</Types>")


def merge_workbook_rels(rels: str, template_rels: str) -> str:
    if "cellimages.xml" in rels:
        return rels
    m = re.search(
        r'<Relationship[^>]+Target="cellimages\.xml"[^>]*/>',
        template_rels,
    )
    if not m:
        raise RuntimeError("Template workbook.xml.rels missing cellimages relationship")
    rel_line = m.group(0)
    if rel_line in rels:
        return rels
    # Avoid Id collision: re-id if necessary
    ids = {int(x) for x in re.findall(r'Id="rId(\d+)"', rels)}
    tid = re.search(r'Id="(rId\d+)"', rel_line)
    if tid and tid.group(1) in rels:
        next_id = max(ids) + 1 if ids else 1
        rel_line = re.sub(r'Id="rId\d+"', f'Id="rId{next_id}"', rel_line)
    return rels.replace("</Relationships>", rel_line + "</Relationships>")


def reinject_media_from_template(target: Path, template: Path) -> bool:
    if not needs_media_reinject(target):
        return False

    with zipfile.ZipFile(target, "r") as zt:
        entries = {name: zt.read(name) for name in zt.namelist()}

    with zipfile.ZipFile(template, "r") as zs:
        template_rels = zs.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        for name in zs.namelist():
            if name.startswith("xl/media/") and not name.endswith("/"):
                entries[name] = zs.read(name)
            elif name in MEDIA_PARTS:
                entries[name] = zs.read(name)

    ct = entries["[Content_Types].xml"].decode("utf-8")
    if 'Extension="png"' not in ct:
        ct = ct.replace(
            "<Default Extension=\"rels\"",
            '<Default Extension="png" ContentType="image/png"/><Default Extension="rels"',
        )
    ct = merge_xml_snippet(
        ct,
        '<Override PartName="/xl/cellimages.xml" ContentType="application/vnd.wps-officedocument.cellimage+xml"/>',
    )
    entries["[Content_Types].xml"] = ct.encode("utf-8")

    rels = entries["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rels = merge_workbook_rels(rels, template_rels)
    entries["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    tmp = target.with_suffix(".mediafix.xlsx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zw:
        for name, data in entries.items():
            zw.writestr(name, data)
    tmp.replace(target)
    return True


def apply_row15_fixed(ws) -> None:
    for ref, val in {
        "I15": 16,
        "J15": 16,
        "K15": 16,
        "L15": 16,
        "M15": 6302539090,
        "N15": "聚酯纤维/polyester",
        "O15": "餐桌",
        "P15": "无",
        "Q15": "无",
        "R15": "无",
        "H15": "USD",
        "D15": "桌布",
        "E15": "table cloth",
    }.items():
        ws[ref] = val


def fix_one_file(dest: Path, template_c2: str) -> dict:
    wh = warehouse_from_name(dest.name)
    info = WAREHOUSES[wh]
    per_carton = math.floor(info["strips"] / info["cartons"])

    shutil.copy2(TEMPLATE, dest)

    wb = openpyxl.load_workbook(dest)
    ws = wb.active

    ws["C2"] = template_c2
    ws["E3"] = wh
    ws["E7"] = None
    ws["E8"] = info["addr"]
    for clear_ref in ("E10", "E11", "E12"):
        ws[clear_ref] = None

    ws["A15"] = info["fba"]
    ws["B15"] = info["track"]
    ws["C15"] = info["cartons"]
    ws["F15"] = per_carton
    ws["G15"] = info["unit"]
    apply_row15_fixed(ws)

    s15 = ws["S15"].value
    if not (isinstance(s15, str) and "DISPIMG" in s15.upper()):
        ws["S15"] = DISPIMG_FORMULA
    ws["T15"] = C04_SALES_URL

    wb.save(dest)
    wb.close()

    reinjected = reinject_media_from_template(dest, TEMPLATE)

    wb2 = openpyxl.load_workbook(dest, data_only=False)
    ws2 = wb2.active
    dump = {ref: cell_str(ws2[ref].value) for ref in KEY_CELLS}
    wb2.close()
    media_n, media_list = count_media(dest)
    dump["_warehouse"] = wh
    dump["_media_count"] = str(media_n)
    dump["_media_files"] = ", ".join(media_list) if media_list else "(none)"
    dump["_media_reinjected"] = str(reinjected)
    return dump


def write_note(processed: list[tuple[Path, dict]], template_c2: str) -> None:
    NOTE_PATH.parent.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M %z")
    lines = [
        "# C26079-A freight forwarder xlsx fix",
        "",
        f"- Timestamp: {ts}",
        f"- Template: {TEMPLATE}",
        f"- Template C2 preserved: {template_c2} (never overwritten with tracking)",
        "",
        "## Corrections applied (freight table only)",
        "",
        "- Reset each file from freight template via shutil.copy2, then set allowed cells only.",
        "- E3 warehouse code; E7 cleared; E8 = warehouse code + address; E10/E11/E12 cleared.",
        "- A15 FBA id; B15 tracking; C15 cartons; F15 = floor(strips/cartons); G15 per-warehouse unit price.",
        "- Restored fixed row-15 customs/product fields (I–R, M HS, N material, O use, H USD).",
        "- S15 DISPIMG formula + re-injected xl/media / cellimages parts from template when openpyxl stripped them.",
        "- Left C3/C5 and 客户单号 C2 unchanged from template rules; T15 = C04 sales URL (https://www.amazon.com/dp/B0CCRKDW1K?th=1).",
        "",
        "## Files",
        "",
    ]
    for path, dump in processed:
        lines.append(f"- {path.name} ({dump['_warehouse']}): media={dump['_media_count']}, reinjected={dump['_media_reinjected']}")
    NOTE_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    if not TEMPLATE.is_file():
        print(f"ERROR: template missing: {TEMPLATE}")
        return 1

    freight_dir = find_freight_dir()
    template_c2 = read_template_c2()

    targets = sorted(
        p
        for p in freight_dir.glob("*.xlsx")
        if not p.name.startswith("~$") and "C26079-A" in p.name
    )
    if len(targets) != 5:
        print(f"ERROR: expected 5 freight files, found {len(targets)} in {freight_dir}")
        for p in targets:
            print(" ", p.name)
        return 1

    processed: list[tuple[Path, dict]] = []
    for path in targets:
        dump = fix_one_file(path, template_c2)
        processed.append((path, dump))

    write_note(processed, template_c2)

    print("C26079-A FREIGHT FIX VERIFICATION")
    print(f"Template C2: {template_c2!r}")
    print(f"Freight dir: {freight_dir}")
    print(f"Note: {NOTE_PATH}")
    print()

    for path, dump in processed:
        print("=" * 90)
        print(path.name)
        print(f"  warehouse: {dump['_warehouse']}")
        print(f"  xl/media count: {dump['_media_count']}  reinjected: {dump['_media_reinjected']}")
        print(f"  media files: {dump['_media_files']}")
        for ref in KEY_CELLS:
            print(f"  {ref}: {dump[ref]}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
