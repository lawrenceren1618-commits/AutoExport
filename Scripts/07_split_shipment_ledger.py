# -*- coding: utf-8 -*-
"""拆单发货比对账本：计划 / 已发 / 剩余。

Usage:
  python Scripts/07_split_shipment_ledger.py init --project "<Shipments/...>"
  python Scripts/07_split_shipment_ledger.py status --project "<...>"
  python Scripts/07_split_shipment_ledger.py compare --project "<...>"
  python Scripts/07_split_shipment_ledger.py add-batch --project "<...>" --batch-json batch.json

batch.json example:
{
  "split_leg_id": "C26079-A",
  "batch_no": 1,
  "ship_date": "2026-07-31",
  "outbound_short": "C04-2616",
  "note": "第一批",
  "shipped_strips": 10000,
  "shipped_cartons": 200,
  "by_size_color": [
    {"size": "60X102inch", "color": "白", "qty_strips": 2000, "qty_cartons": 50}
  ],
  "by_fnsku": [
    {"size": "60X102inch", "color": "白", "strips_per_pack": 10, "fnsku": "X003WN3V5J",
     "qty_strips": 2000, "qty_cartons": 50}
  ],
  "warehouses": [{"warehouse_code": "HEA2", "cartons": 200, "strips": 10000}]
}
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]


def _now() -> str:
    return datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%dT%H:%M:%S%z")


def _norm_size(s: str) -> str:
    if not s:
        return ""
    t = str(s).strip().replace("＊", "*").replace("×", "X").replace("x", "X")
    t = t.replace("英寸", "inch").replace(" ", "")
    # restore space in round
    if "round" in t.lower():
        t = re.sub(r"(?i)90\s*inch\s*round", "90 inch round", t.replace("90inchround", "90 inch round"))
        return "90 inch round" if "90" in t and "round" in t.lower() else t
    t = t.replace("*", "X")
    # common pads
    for a, b in [
        ("60X102inch", "60X102inch"),
        ("60X126inch", "60X126inch"),
        ("70X120inch", "70X120inch"),
        ("60X84inch", "60X84inch"),
    ]:
        if a.lower() == t.lower() or t.upper() == a.upper():
            return a
    if re.fullmatch(r"60X102INCH", t, re.I):
        return "60X102inch"
    if re.fullmatch(r"60X126INCH", t, re.I):
        return "60X126inch"
    if re.fullmatch(r"70X120INCH", t, re.I):
        return "70X120inch"
    return t


def _norm_color(c: str) -> str:
    if not c:
        return ""
    c = str(c).strip()
    aliases = {
        "白": "白",
        "白色": "白",
        "黑": "黑",
        "黑色": "黑",
        "蓝": "蓝",
        "蓝色": "蓝",
        "米白": "米白",
        "米白色": "米白",
        "象牙": "米白",
        "象牙色": "米白",
        "象牙白": "米白",
        "象牙白色": "米白",
        "Ivory": "米白",
        "ivory": "米白",
    }
    return aliases.get(c, c)


def _ledger_dir(project: Path) -> Path:
    d = project / "拆单比对"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _ledger_path(project: Path) -> Path:
    return _ledger_dir(project) / "split_ledger.json"


def order_stem_from_code(order_code: str) -> str:
    """EF-C26079-C04-36460 -> C26079."""
    if not order_code:
        return ""
    m = re.search(r"EF-(C\d{5})-", str(order_code), re.I)
    if m:
        return m.group(1).upper()
    m = re.search(r"C\d{5}", str(order_code), re.I)
    if m:
        return m.group(0).upper()
    return ""


def _letter_from_index(n: int) -> str:
    """1 -> A, 26 -> Z, 27 -> AA."""
    if n <= 0:
        raise ValueError(n)
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _index_from_letter(letter: str) -> int:
    letter = str(letter or "").upper()
    n = 0
    for ch in letter:
        if not ("A" <= ch <= "Z"):
            continue
        n = n * 26 + (ord(ch) - 64)
    return n


def split_leg_id_from_stem(stem: str, letter: str) -> str:
    return f"{stem}-{str(letter).upper()}"


def _parse_split_leg_id(leg_id: str) -> tuple[str, str] | None:
    if not leg_id:
        return None
    m = re.fullmatch(r"(C\d{5})-([A-Z]+)", str(leg_id).strip(), re.I)
    if not m:
        return None
    return m.group(1).upper(), m.group(2).upper()


def _batch_no_from_split_leg_id(leg_id: str) -> int | None:
    parsed = _parse_split_leg_id(leg_id)
    if not parsed:
        return None
    return _index_from_letter(parsed[1])


def _highest_used_leg_index(batches: list[dict]) -> int:
    best = 0
    for b in batches or []:
        leg = b.get("split_leg_id")
        if leg:
            parsed = _parse_split_leg_id(leg)
            if parsed:
                best = max(best, _index_from_letter(parsed[1]))
                continue
        try:
            best = max(best, int(b.get("batch_no") or 0))
        except (TypeError, ValueError):
            pass
    return best


def _update_split_leg_state(ledger: dict) -> None:
    """Set split_leg_next and completion note after remaining is computed."""
    stem = ledger.get("order_stem") or order_stem_from_code(ledger.get("order_code", ""))
    if stem:
        ledger["order_stem"] = stem
    if not ledger.get("split_naming"):
        ledger["split_naming"] = "order_stem_letter"

    rem_strips = int(ledger.get("remaining", {}).get("qty_strips") or 0)
    rem_cartons = int(ledger.get("remaining", {}).get("qty_cartons") or 0)
    batches = ledger.get("batches") or []
    warnings = ledger.get("warnings") or []
    cartons_ok = rem_cartons >= 0

    ledger.pop("split_complete_note", None)

    if rem_strips == 0 and cartons_ok and not warnings:
        ledger["status"] = "split_complete"
        ledger["split_leg_next"] = None
        if len(batches) == 1:
            leg = batches[0].get("split_leg_id") or ""
            if leg.upper().endswith("-A"):
                ledger["split_complete_note"] = "A filled full order → no B/C"
        return

    if batches:
        ledger["status"] = "split_in_progress"
    else:
        ledger["status"] = "split_pending_batch1"

    if not stem:
        return

    hi = _highest_used_leg_index(batches)
    if hi <= 0:
        ledger["split_leg_next"] = split_leg_id_from_stem(stem, "A")
    else:
        ledger["split_leg_next"] = split_leg_id_from_stem(stem, _letter_from_index(hi + 1))


def _load_ledger(project: Path) -> dict:
    p = _ledger_path(project)
    if not p.exists():
        raise FileNotFoundError(f"账本不存在，请先 init: {p}")
    return json.loads(p.read_text(encoding="utf-8"))


def _save_ledger(project: Path, data: dict) -> Path:
    p = _ledger_path(project)
    data["updated_at"] = _now()
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


def _find_order_xls(project: Path) -> Path:
    order_dir = project / "订单表"
    if not order_dir.is_dir():
        raise FileNotFoundError(f"无订单表目录: {order_dir}")
    cands = list(order_dir.glob("*.xls")) + list(order_dir.glob("*.xlsx"))
    if not cands:
        raise FileNotFoundError(f"订单表目录无 xls/xlsx: {order_dir}")
    return cands[0]


def _parse_order_xls(path: Path) -> tuple[list[dict], list[dict], int, int]:
    """Return upper lines, lower lines, total strips, total cartons."""
    if path.suffix.lower() == ".xls":
        if xlrd is None:
            raise RuntimeError("需要 xlrd 读取 .xls")
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)

        def cell(r, c):
            if r >= sheet.nrows or c >= sheet.ncols:
                return ""
            v = sheet.cell_value(r, c)
            if isinstance(v, float) and v == int(v):
                return str(int(v))
            return str(v).strip() if v is not None else ""

        nrows = sheet.nrows
    else:
        if openpyxl is None:
            raise RuntimeError("需要 openpyxl 读取 .xlsx")
        wb = openpyxl.load_workbook(str(path), data_only=True)
        sheet = wb.active

        def cell(r, c):
            v = sheet.cell(r + 1, c + 1).value
            if v is None:
                return ""
            if isinstance(v, float) and v == int(v):
                return str(int(v))
            return str(v).strip()

        nrows = sheet.max_row or 0

    # find lower header row (条数/包)
    lower_header = None
    for r in range(nrows):
        rowtxt = " ".join(cell(r, c) for c in range(12))
        if "条数/包" in rowtxt and ("产品标签" in rowtxt or "FNSKU" in rowtxt.upper() or "产品标" in rowtxt):
            lower_header = r
            break
        if "条数/包" in rowtxt and "总箱数" in rowtxt:
            lower_header = r
            break

    if lower_header is None:
        raise RuntimeError("未找到下半装箱明细表头（条数/包）")

    # upper: from first data after 尺寸规格 header to before 合计 / 装箱
    upper_header = None
    for r in range(lower_header):
        if cell(r, 0) == "订单编号" or (cell(r, 3) in ("尺寸规格", "尺寸项") and cell(r, 4) in ("颜色", "颜 色")):
            upper_header = r
            break
        rowtxt = "".join(cell(r, c) for c in range(7))
        if "尺寸规格" in rowtxt and "颜色" in rowtxt and "数量" in rowtxt:
            upper_header = r
            break

    if upper_header is None:
        upper_header = 7  # fallback R8

    upper: list[dict] = []
    last_size = ""
    for r in range(upper_header + 1, lower_header):
        c0 = cell(r, 0)
        if "合计" in c0.replace(" ", "") or "制单" in c0 or "装箱" in c0:
            break
        size = cell(r, 3) or last_size
        color = cell(r, 4)
        qty_s = cell(r, 5)
        if not color or not qty_s:
            continue
        try:
            qty = int(float(qty_s))
        except ValueError:
            continue
        if size:
            last_size = size
        upper.append(
            {
                "size": _norm_size(size or last_size),
                "color": _norm_color(color),
                "qty_strips": qty,
                "qty_cartons": None,
            }
        )

    lower: list[dict] = []
    total_cartons = 0
    total_strips_lower = 0
    for r in range(lower_header + 1, nrows):
        size = cell(r, 0)
        if not size or "合计" in size:
            # try total row
            if "合计" in size or cell(r, 0).replace(" ", "") == "合计":
                tc = cell(r, 5)
                ts = cell(r, 6)
                try:
                    if tc:
                        total_cartons = int(float(tc))
                    if ts:
                        total_strips_lower = int(float(ts))
                except ValueError:
                    pass
            break
        color = cell(r, 1)
        spp = cell(r, 2)
        fnsku = cell(r, 3)
        packs = cell(r, 4)
        cartons = cell(r, 5)
        strips = cell(r, 6)
        try:
            spp_i = int(float(spp)) if spp else None
            cartons_i = int(float(cartons)) if cartons else 0
            strips_i = int(float(strips)) if strips else 0
        except ValueError:
            continue
        if not color or strips_i <= 0:
            continue
        lower.append(
            {
                "size": _norm_size(size),
                "color": _norm_color(color),
                "strips_per_pack": spp_i,
                "fnsku": fnsku,
                "packs_per_carton": int(float(packs)) if packs else None,
                "qty_strips": strips_i,
                "qty_cartons": cartons_i,
            }
        )
        total_cartons += cartons_i
        total_strips_lower += strips_i

    total_strips_upper = sum(x["qty_strips"] for x in upper)
    total_strips = total_strips_lower or total_strips_upper
    if total_cartons == 0:
        total_cartons = sum(x.get("qty_cartons") or 0 for x in lower)

    return upper, lower, total_strips, total_cartons


def _zero_remaining_from_plan(plan: dict) -> dict:
    return {
        "qty_strips": plan["qty_strips"],
        "qty_cartons": plan["qty_cartons"],
        "by_size_color": [
            {
                "size": x["size"],
                "color": x["color"],
                "qty_strips": x["qty_strips"],
                "qty_cartons": x.get("qty_cartons"),
            }
            for x in plan["by_size_color"]
        ],
        "by_fnsku": [
            {
                "size": x["size"],
                "color": x["color"],
                "strips_per_pack": x.get("strips_per_pack"),
                "fnsku": x.get("fnsku"),
                "qty_strips": x["qty_strips"],
                "qty_cartons": x.get("qty_cartons"),
            }
            for x in plan["by_fnsku"]
        ],
    }


def _key_sc(x: dict) -> tuple:
    return (_norm_size(x.get("size", "")), _norm_color(x.get("color", "")))


def _key_fn(x: dict) -> tuple:
    return (
        _norm_size(x.get("size", "")),
        _norm_color(x.get("color", "")),
        x.get("strips_per_pack"),
        (x.get("fnsku") or "").strip(),
    )


def _recompute(ledger: dict) -> list[str]:
    """Update shipped_total + remaining; return list of imbalance warnings."""
    plan = ledger["plan"]
    warnings: list[str] = []

    shipped_strips = 0
    shipped_cartons = 0
    sc_ship: dict[tuple, dict] = {}
    fn_ship: dict[tuple, dict] = {}

    for b in ledger.get("batches", []):
        shipped_strips += int(b.get("shipped_strips") or 0)
        shipped_cartons += int(b.get("shipped_cartons") or 0)
        for line in b.get("by_size_color") or []:
            k = _key_sc(line)
            sc_ship.setdefault(k, {"size": k[0], "color": k[1], "qty_strips": 0, "qty_cartons": 0})
            sc_ship[k]["qty_strips"] += int(line.get("qty_strips") or 0)
            sc_ship[k]["qty_cartons"] += int(line.get("qty_cartons") or 0)
        for line in b.get("by_fnsku") or []:
            k = _key_fn(line)
            fn_ship.setdefault(
                k,
                {
                    "size": k[0],
                    "color": k[1],
                    "strips_per_pack": k[2],
                    "fnsku": k[3],
                    "qty_strips": 0,
                    "qty_cartons": 0,
                },
            )
            fn_ship[k]["qty_strips"] += int(line.get("qty_strips") or 0)
            fn_ship[k]["qty_cartons"] += int(line.get("qty_cartons") or 0)

    ledger["shipped_total"] = {
        "qty_strips": shipped_strips,
        "qty_cartons": shipped_cartons,
        "by_size_color": list(sc_ship.values()),
        "by_fnsku": list(fn_ship.values()),
    }

    rem_sc = []
    for p in plan["by_size_color"]:
        k = _key_sc(p)
        got = sc_ship.get(k, {})
        left_s = int(p["qty_strips"]) - int(got.get("qty_strips") or 0)
        left_c = None
        if p.get("qty_cartons") is not None:
            left_c = int(p["qty_cartons"]) - int(got.get("qty_cartons") or 0)
        if left_s < 0 or (left_c is not None and left_c < 0):
            warnings.append(f"L1超发 {k}: remaining strips={left_s} cartons={left_c}")
        rem_sc.append(
            {
                "size": k[0],
                "color": k[1],
                "qty_strips": left_s,
                "qty_cartons": left_c,
                "plan_strips": p["qty_strips"],
                "shipped_strips": int(got.get("qty_strips") or 0),
            }
        )

    rem_fn = []
    for p in plan["by_fnsku"]:
        k = _key_fn(p)
        got = fn_ship.get(k, {})
        left_s = int(p["qty_strips"]) - int(got.get("qty_strips") or 0)
        left_c = int(p.get("qty_cartons") or 0) - int(got.get("qty_cartons") or 0)
        if left_s < 0 or left_c < 0:
            warnings.append(f"L2超发 {k}: remaining strips={left_s} cartons={left_c}")
        rem_fn.append(
            {
                "size": k[0],
                "color": k[1],
                "strips_per_pack": k[2],
                "fnsku": k[3],
                "qty_strips": left_s,
                "qty_cartons": left_c,
                "plan_strips": p["qty_strips"],
                "shipped_strips": int(got.get("qty_strips") or 0),
                "plan_cartons": p.get("qty_cartons"),
                "shipped_cartons": int(got.get("qty_cartons") or 0),
            }
        )

    rem_strips = int(plan["qty_strips"]) - shipped_strips
    rem_cartons = int(plan["qty_cartons"]) - shipped_cartons
    if rem_strips < 0 or rem_cartons < 0:
        warnings.append(f"L0超发: remaining strips={rem_strips} cartons={rem_cartons}")

    ledger["remaining"] = {
        "qty_strips": rem_strips,
        "qty_cartons": rem_cartons,
        "by_size_color": rem_sc,
        "by_fnsku": rem_fn,
    }

    ledger["balance_ok"] = len(warnings) == 0 and rem_strips >= 0 and rem_cartons >= 0
    ledger["warnings"] = warnings
    _update_split_leg_state(ledger)
    return warnings


def _write_report(project: Path, ledger: dict) -> Path:
    lines = [
        f"# 拆单比对报告 — {ledger.get('order_code')}",
        "",
        f"- 更新: {ledger.get('updated_at')}",
        f"- 状态: `{ledger.get('status')}`",
        f"- 轧平: {'OK' if ledger.get('balance_ok') else 'CHECK WARNINGS'}",
        f"- 拆单命名: `{ledger.get('split_naming', 'order_stem_letter')}`",
        f"- 订单 stem: `{ledger.get('order_stem') or ''}`",
        f"- 下一腿: `{ledger.get('split_leg_next')}`" + (
            "（A 已发满整单计划，无需 B/C）" if ledger.get('split_leg_next') is None and ledger.get('status') == 'split_complete' else ""
        ),
    ]
    if ledger.get("split_complete_note"):
        lines.append(f"- 说明: {ledger['split_complete_note']}")
    lines.extend([
        "",
        "",
        "## L0 整单",
        "",
        "| 项 | 计划 | 已发合计 | 剩余 |",
        "|----|------|----------|------|",
        f"| 条数 | {ledger['plan']['qty_strips']} | {ledger['shipped_total']['qty_strips']} | {ledger['remaining']['qty_strips']} |",
        f"| 箱数 | {ledger['plan']['qty_cartons']} | {ledger['shipped_total']['qty_cartons']} | {ledger['remaining']['qty_cartons']} |",
        "",
        "## 批次",
        "",
    ])
    if not ledger.get("batches"):
        lines.append("_尚无批次；放入 FBA 外箱贴或提供 batch.json 后执行 add-batch。_")
        lines.append("")
    else:
        lines.append("| 腿 ID | 批 | 短号 | 日期 | 条数 | 箱数 | 备注 |")
        lines.append("|-------|----|------|------|------|------|------|")
        for b in ledger["batches"]:
            lines.append(
                f"| {b.get('split_leg_id') or ''} | {b.get('batch_no')} | {b.get('outbound_short') or ''} | {b.get('ship_date') or ''} | "
                f"{b.get('shipped_strips')} | {b.get('shipped_cartons')} | {b.get('note') or ''} |"
            )
        lines.append("")

    lines.extend(
        [
            "## L1 剩余（尺寸+颜色，出货口径）",
            "",
            "| 尺寸 | 颜色 | 计划条 | 已发条 | 剩余条 |",
            "|------|------|--------|--------|--------|",
        ]
    )
    for x in ledger["remaining"]["by_size_color"]:
        if x["qty_strips"] == 0 and x.get("shipped_strips", 0) == 0:
            continue
        lines.append(
            f"| {x['size']} | {x['color']} | {x.get('plan_strips', '')} | {x.get('shipped_strips', 0)} | {x['qty_strips']} |"
        )

    lines.extend(
        [
            "",
            "## L2 剩余（FNSKU）摘要",
            "",
            f"- 仍有剩余的 FNSKU 行数: {sum(1 for x in ledger['remaining']['by_fnsku'] if x['qty_strips'] > 0)}",
            f"- 已发完的 FNSKU 行数: {sum(1 for x in ledger['remaining']['by_fnsku'] if x['qty_strips'] == 0)}",
            "",
        ]
    )
    if ledger.get("warnings"):
        lines.append("## 警告")
        lines.append("")
        for w in ledger["warnings"]:
            lines.append(f"- {w}")
        lines.append("")

    out = _ledger_dir(project) / "_last_compare_report.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def cmd_init(project: Path) -> None:
    order = _find_order_xls(project)
    upper, lower, strips, cartons = _parse_order_xls(order)
    proj_meta = {}
    pj = project / "PROJECT.json"
    if pj.exists():
        proj_meta = json.loads(pj.read_text(encoding="utf-8"))

    plan = {
        "qty_strips": strips,
        "qty_cartons": cartons,
        "by_size_color": upper,
        "by_fnsku": lower,
        "source_order": str(order),
    }
    order_code = proj_meta.get("order_code") or project.name.split()[0]
    order_stem = order_stem_from_code(order_code)
    ledger = {
        "version": "1.0",
        "order_code": order_code,
        "order_stem": order_stem,
        "split_naming": "order_stem_letter",
        "split_leg_next": split_leg_id_from_stem(order_stem, "A") if order_stem else None,
        "project_folder": str(project),
        "split_shipment": True,
        "created_at": _now(),
        "plan": plan,
        "batches": [],
        "shipped_total": {
            "qty_strips": 0,
            "qty_cartons": 0,
            "by_size_color": [],
            "by_fnsku": [],
        },
        "remaining": _zero_remaining_from_plan(plan),
        "status": "split_pending_batch1",
        "balance_ok": True,
        "warnings": [],
        "doc": "Templates/Split_Shipment_Ledger.md",
    }
    # annotate remaining with plan/shipped for report
    _recompute(ledger)
    path = _save_ledger(project, ledger)
    report = _write_report(project, ledger)

    if pj.exists():
        proj_meta["split_shipment"] = True
        proj_meta["split_ledger"] = str(path)
        proj_meta["split_order_stem"] = order_stem
        proj_meta["split_leg_next"] = ledger.get("split_leg_next")
        proj_meta["split_naming"] = f"{order_stem}-A/B/C" if order_stem else ledger.get("split_naming")
        proj_meta["status"] = ledger.get("status", "split_pending_batch1")
        pj.write_text(json.dumps(proj_meta, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"INIT OK: {path}")
    print(f"plan strips={strips} cartons={cartons} L1={len(upper)} L2={len(lower)}")
    print(f"report: {report}")


def cmd_status(project: Path) -> None:
    ledger = _load_ledger(project)
    _recompute(ledger)
    _save_ledger(project, ledger)
    report = _write_report(project, ledger)
    print(json.dumps(
        {
            "status": ledger["status"],
            "order_stem": ledger.get("order_stem"),
            "split_leg_next": ledger.get("split_leg_next"),
            "split_naming": ledger.get("split_naming"),
            "split_complete_note": ledger.get("split_complete_note"),
            "plan_strips": ledger["plan"]["qty_strips"],
            "shipped_strips": ledger["shipped_total"]["qty_strips"],
            "remaining_strips": ledger["remaining"]["qty_strips"],
            "plan_cartons": ledger["plan"]["qty_cartons"],
            "shipped_cartons": ledger["shipped_total"]["qty_cartons"],
            "remaining_cartons": ledger["remaining"]["qty_cartons"],
            "batches": len(ledger.get("batches") or []),
            "balance_ok": ledger["balance_ok"],
            "warnings": ledger.get("warnings"),
            "report": str(report),
        },
        ensure_ascii=False,
        indent=2,
    ))


def cmd_compare(project: Path) -> None:
    cmd_status(project)


def cmd_add_batch(project: Path, batch_json: Path) -> None:
    ledger = _load_ledger(project)
    batch = json.loads(Path(batch_json).read_text(encoding="utf-8"))
    stem = ledger.get("order_stem") or order_stem_from_code(ledger.get("order_code", ""))
    if stem:
        ledger["order_stem"] = stem
    leg = batch.get("split_leg_id") or ledger.get("split_leg_next")
    if not leg and stem:
        leg = split_leg_id_from_stem(stem, "A")
    if leg:
        batch["split_leg_id"] = leg
        derived = _batch_no_from_split_leg_id(leg)
        if derived is not None:
            batch["batch_no"] = derived
    if "batch_no" not in batch:
        batch["batch_no"] = len(ledger.get("batches") or []) + 1
    # normalize lines
    for line in batch.get("by_size_color") or []:
        line["size"] = _norm_size(line.get("size", ""))
        line["color"] = _norm_color(line.get("color", ""))
    for line in batch.get("by_fnsku") or []:
        line["size"] = _norm_size(line.get("size", ""))
        line["color"] = _norm_color(line.get("color", ""))

    # replace same batch_no if re-run
    def _same_batch(a: dict, b: dict) -> bool:
        if a.get("split_leg_id") and b.get("split_leg_id"):
            return str(a["split_leg_id"]).upper() == str(b["split_leg_id"]).upper()
        return int(a.get("batch_no") or 0) == int(b.get("batch_no") or 0)

    batches = [b for b in ledger.get("batches") or [] if not _same_batch(b, batch)]
    batches.append(batch)
    batches.sort(key=lambda b: int(b.get("batch_no") or 0))
    ledger["batches"] = batches

    warnings = _recompute(ledger)
    path = _save_ledger(project, ledger)
    report = _write_report(project, ledger)

    # sidecar copy of batch
    leg_id = batch.get("split_leg_id") or f"batch_{batch['batch_no']}"
    side = _ledger_dir(project) / f"{leg_id}_manifest.json"
    side.write_text(json.dumps(batch, ensure_ascii=False, indent=2), encoding="utf-8")

    pj = project / "PROJECT.json"
    if pj.exists():
        meta = json.loads(pj.read_text(encoding="utf-8"))
        meta["split_shipment"] = True
        meta["split_batch_latest"] = batch["batch_no"]
        meta["split_leg_latest"] = batch.get("split_leg_id")
        meta["split_leg_next"] = ledger.get("split_leg_next")
        meta["split_order_stem"] = ledger.get("order_stem")
        if ledger.get("order_stem"):
            meta["split_naming"] = f"{ledger['order_stem']}-A/B/C"
        meta["remaining_strips"] = ledger["remaining"]["qty_strips"]
        meta["remaining_cartons"] = ledger["remaining"]["qty_cartons"]
        meta["status"] = ledger["status"]
        pj.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"ADD-BATCH OK: split_leg_id={batch.get('split_leg_id')} batch_no={batch['batch_no']}")
    print(f"ledger: {path}")
    print(f"remaining strips={ledger['remaining']['qty_strips']} cartons={ledger['remaining']['qty_cartons']}")
    if warnings:
        print("WARNINGS:")
        for w in warnings:
            print(" -", w)
    print(f"report: {report}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="拆单发货比对账本")
    ap.add_argument("command", choices=["init", "status", "compare", "add-batch"])
    ap.add_argument("--project", required=True, help="Shipments 项目文件夹路径")
    ap.add_argument("--batch-json", help="add-batch 用的本批 JSON")
    args = ap.parse_args(argv)

    project = Path(args.project)
    if not project.is_dir():
        print(f"项目目录不存在: {project}", file=sys.stderr)
        return 1

    if args.command == "init":
        cmd_init(project)
    elif args.command in ("status", "compare"):
        cmd_status(project)
    elif args.command == "add-batch":
        if not args.batch_json:
            print("add-batch 需要 --batch-json", file=sys.stderr)
            return 1
        cmd_add_batch(project, Path(args.batch_json))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
