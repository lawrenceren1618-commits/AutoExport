# -*- coding: utf-8 -*-
"""Ingest Amazon shipment tracking from screenshots / structured JSON.

Inbox: Shipments/{project}/货代入库贴下单表/_tracking_inbox/
Writes tracking_id into freight xlsx B15 (does NOT touch customer no C2=111111111).

Usage:
  # Apply known mapping JSON (from OCR or manual after screenshot upload)
  python Scripts/14_tracking_screenshot_ingest.py --project "Shipments/..." --mapping drafts/xxx_tracking.json --apply

  # List inbox only
  python Scripts/14_tracking_screenshot_ingest.py --project "Shipments/..." --list-inbox
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
TZ = timezone(timedelta(hours=8))

FBA_RE = re.compile(r"\b(FBA[0-9A-Z]{8,12})\b", re.I)
WH_RE = re.compile(r"\b([A-Z]{3}\d)\b")
# Amazon tracking-ish tokens near 追踪
TRACK_RE = re.compile(
    r"(?:货件追踪编号|追踪编号|Tracking\s*(?:ID|number)?)\s*[:：]?\s*([A-Z0-9]{6,14})",
    re.I,
)


def now() -> str:
    return datetime.now(TZ).strftime("%Y-%m-%dT%H:%M:%S%z")


def find_project(arg: str) -> Path:
    p = Path(arg)
    if not p.is_absolute():
        p = ROOT / p
    if p.is_dir():
        return p
    for cand in (ROOT / "1_Input_Raw/Shipments").iterdir():
        if arg in cand.name or "C26080" in cand.name:
            return cand
    raise SystemExit(f"project not found: {arg}")


def inbox_dir(project: Path) -> Path:
    d = project / "货代入库贴下单表" / "_tracking_inbox"
    d.mkdir(parents=True, exist_ok=True)
    return d


def list_inbox(project: Path) -> list[Path]:
    d = inbox_dir(project)
    return sorted(
        p
        for p in d.iterdir()
        if p.is_file() and p.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".json"}
    )


def try_ocr_text(image: Path) -> str:
    """Best-effort OCR; returns empty if libs unavailable."""
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore

        return pytesseract.image_to_string(Image.open(image), lang="chi_sim+eng") or ""
    except Exception:
        return ""


def parse_text_blocks(text: str) -> list[dict]:
    """Extract warehouse / FBA / tracking triples from OCR or pasted text."""
    rows: list[dict] = []
    # split by shipment-ish chunks
    chunks = re.split(r"(?:货件\s*\d+|Shipment\s*\d+)", text, flags=re.I)
    if len(chunks) <= 1:
        chunks = [text]
    for ch in chunks:
        whs = WH_RE.findall(ch)
        fbas = [x.upper() for x in FBA_RE.findall(ch)]
        tracks = TRACK_RE.findall(ch)
        if not tracks:
            # fallback: line after 追踪
            for m in re.finditer(r"追踪[^\n]{0,20}\n\s*([A-Z0-9]{6,14})", ch, re.I):
                tracks.append(m.group(1))
        if not (whs or fbas or tracks):
            continue
        rows.append(
            {
                "warehouse_code": whs[0] if whs else "",
                "fba_shipment_id": fbas[0] if fbas else "",
                "tracking_id": tracks[0].upper() if tracks else "",
                "raw_excerpt": ch[:200].replace("\n", " "),
            }
        )
    return rows


def load_mapping(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict) and "rows" in data:
        return data["rows"]
    if isinstance(data, list):
        return data
    raise SystemExit("mapping must be list or {rows:[...]}")


def freight_files(project: Path) -> list[Path]:
    d = project / "货代入库贴下单表"
    return sorted(d.glob("货代*.xlsx"))


def apply_tracking(project: Path, rows: list[dict]) -> dict:
    by_wh = {str(r.get("warehouse_code") or "").upper(): r for r in rows if r.get("warehouse_code")}
    by_fba = {
        str(r.get("fba_shipment_id") or "").upper(): r
        for r in rows
        if r.get("fba_shipment_id")
    }
    updated = []
    missing = []
    for fpath in freight_files(project):
        wb = openpyxl.load_workbook(fpath)
        ws = wb["下单模板"]
        wh = str(ws["E3"].value or "").upper()
        fba = str(ws["A15"].value or "").upper()
        row = by_wh.get(wh) or by_fba.get(fba)
        if not row or not row.get("tracking_id"):
            missing.append({"file": fpath.name, "warehouse_code": wh, "fba": fba})
            wb.close()
            continue
        track = str(row["tracking_id"]).strip().upper()
        ws["B15"] = track
        # never touch C2 customer no
        wb.save(fpath)
        wb.close()
        updated.append(
            {
                "file": fpath.name,
                "warehouse_code": wh,
                "fba_shipment_id": fba,
                "tracking_id": track,
            }
        )
    return {
        "generated_at": now(),
        "updated": updated,
        "missing": missing,
        "rule": "Templates/Shipment_Tracking_Screenshot_Inbox.md",
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True)
    ap.add_argument("--mapping", default="", help="JSON mapping rows")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--list-inbox", action="store_true")
    ap.add_argument("--ocr-inbox", action="store_true", help="Try OCR on inbox images → draft JSON")
    args = ap.parse_args()
    project = find_project(args.project)
    inbox = inbox_dir(project)

    if args.list_inbox:
        files = list_inbox(project)
        print(json.dumps({"inbox": str(inbox), "files": [p.name for p in files]}, ensure_ascii=False, indent=2))
        return 0

    rows: list[dict] = []
    if args.mapping:
        rows = load_mapping(Path(args.mapping) if Path(args.mapping).is_file() else ROOT / args.mapping)
    elif args.ocr_inbox:
        for img in list_inbox(project):
            if img.suffix.lower() == ".json":
                continue
            text = try_ocr_text(img)
            parsed = parse_text_blocks(text)
            for p in parsed:
                p["source_image"] = img.name
            rows.extend(parsed)
        draft = (
            ROOT
            / "2_Processed_AI/drafts"
            / f"{project.name}_tracking_ocr_draft.json"
        )
        draft.write_text(
            json.dumps({"generated_at": now(), "rows": rows, "note": "review before --apply"}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(json.dumps({"draft": str(draft), "rows": rows}, ensure_ascii=False, indent=2))
        if not args.apply:
            return 0
    else:
        print("provide --mapping JSON or --ocr-inbox", file=sys.stderr)
        return 1

    if not args.apply:
        print(json.dumps({"dry_run_rows": rows}, ensure_ascii=False, indent=2))
        return 0

    result = apply_tracking(project, rows)
    out = ROOT / "2_Processed_AI/drafts" / "C26080_tracking_apply.json"
    # generic name from stem
    stem = re.search(r"(C\d{5})", project.name)
    if stem:
        out = ROOT / "2_Processed_AI/drafts" / f"{stem.group(1)}_tracking_apply.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    print("wrote", out)
    return 0 if not result["missing"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
