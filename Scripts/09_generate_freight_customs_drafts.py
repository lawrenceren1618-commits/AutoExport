# -*- coding: utf-8 -*-
"""Generate C26079-A freight + customs drafts (one xlsx per warehouse).

Writes under project 货代入库贴下单表/ and 报关资料表/, and mirrors JSON summary
to 2_Processed_AI/drafts/. Does NOT write 3_Output_Final (Gate still required).

Usage:
  python Scripts/09_generate_freight_customs_drafts.py --project "<Shipments/...>"
"""
from __future__ import annotations

import argparse
import json
import math
import random
import re
import shutil
import sys
import zipfile
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Scripts"))
from contract_protocol import allocate_protocol_no  # noqa: E402
from m1.warehouse_address import get_address, load_addresses, placeholder as wh_placeholder  # noqa: E402

FREIGHT_TEMPLATE = ROOT / (
    "1_Input_Raw/Shipments/EF-C26040-C04-8490_C04_20260510_8490/"
    "货代入库贴下单表/货代入库贴数据2026.5.10 HEA2 248箱.xlsx"
)
# 样例固定：客户单号不得改为追踪码/FBA；箱重尺默认 16；图片用样例 DISPIMG+media
FREIGHT_CUSTOMER_NO = "111111111"
FREIGHT_CARTON_DIM_DEFAULT = 16
FREIGHT_DISPIMG = '=DISPIMG("ID_3D20B9B273A54ACBA2D957C7EE1766FF",1)'
FREIGHT_SALES_LINK = {
    "C04": "https://www.amazon.com/dp/B0CCRKDW1K?th=1",
    "C07": "https://www.amazon.com/gp/product/B0FNVMX537",
}
FREIGHT_MEDIA_PARTS = (
    "xl/cellimages.xml",
    "xl/_rels/cellimages.xml.rels",
)
CUSTOMS_TEMPLATE = ROOT / (
    "1_Input_Raw/Shipments/EF-C26040-C04-8490_C04_20260510_8490/"
    "_来源归档/畅选+波丽丝 2026.5.10发货 248箱 市场采购数据.xlsx"
)
MASTER = ROOT / "1_Input_Raw/00_Master_Reference/2026.7.20桌布当前产品总表.xlsx"
CANONICAL = ROOT / "Scripts" / "canonical_schema.json"
WH_ADDRESS: dict = {}  # filled in main() from amazon_warehouse_addresses.json

# Legacy hardcoded suffixes kept only as documentation of C26079-A history.
# New generations MUST use contract_protocol.allocate_protocol_no (registry).
PROTOCOL_SUFFIX = {
    "AVP1": "801",
    "CLT2": "802",
    "GYR2": "803",
    "LGB8": "804",
    "MDW2": "805",
}


def _now() -> str:
    return datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%dT%H:%M:%S%z")


def _parse_lwh(text: str) -> tuple[float, float, float]:
    if not text:
        return (57.0, 40.0, 33.0)
    m = re.search(r"(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)\s*[*×xX]\s*(\d+(?:\.\d+)?)", str(text))
    if m:
        return (float(m.group(1)), float(m.group(2)), float(m.group(3)))
    return (57.0, 40.0, 33.0)


def _load_size_color_maps() -> tuple[dict, dict]:
    short_to_full: dict[str, str] = {}
    color_to_canon: dict[str, str] = {}
    if CANONICAL.is_file():
        schema = json.loads(CANONICAL.read_text(encoding="utf-8"))
        short_to_full = dict(
            ((schema.get("size_shortcode_map") or {}).get("master_short_to_order_full") or {})
        )
        for canon, aliases in ((schema.get("color_aliases") or {}).get("canonical") or {}).items():
            color_to_canon[str(canon).strip()] = str(canon)
            for a in aliases or []:
                color_to_canon[str(a).strip()] = str(canon)
                color_to_canon[str(a).strip().lower()] = str(canon)
    return short_to_full, color_to_canon


def _norm_size(raw, short_to_full: dict) -> str:
    s = str(raw or "").strip()
    if not s:
        return ""
    if s in short_to_full:
        return str(short_to_full[s])
    compact = re.sub(r"[\s*×x]", "", s, flags=re.I).lower()
    for full in short_to_full.values():
        f = str(full)
        if s.lower() == f.lower():
            return f
        if compact == re.sub(r"[\s*×x]", "", f, flags=re.I).lower():
            return f
    return s


def _norm_color(raw, color_to_canon: dict) -> str:
    s = str(raw or "").strip()
    if not s:
        return ""
    return color_to_canon.get(s) or color_to_canon.get(s.lower()) or s


def _eval_master_price_formula(raw, packs_per_carton) -> float | None:
    """data_only 无缓存时，总表 K 列常见式 =39.99*0.55*H{n} → 算每箱1039。"""
    if isinstance(raw, (int, float)):
        return float(raw)
    if not isinstance(raw, str) or not raw.startswith("="):
        return None
    expr = raw[1:].strip().upper()
    # 仅允许数字、*、/、+、-、括号、及 H 列引用
    if not re.fullmatch(r"[0-9H.\+\-\*/\(\)\s]+", expr):
        return None
    if packs_per_carton in (None, ""):
        return None
    expr = re.sub(r"H\d+", str(float(packs_per_carton)), expr)
    try:
        return float(eval(expr, {"__builtins__": {}}, {}))  # noqa: S307 — 已白名单
    except Exception:
        return None


def load_master_index() -> dict:
    """FNSKU index + spec index (条数/包+尺寸+颜色) for C07 inherit of C04 1039."""
    short_to_full, color_to_canon = _load_size_color_maps()
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb["Sheet1"]
    wb_f = openpyxl.load_workbook(MASTER, data_only=False)
    ws_f = wb_f["Sheet1"]
    spp = None
    size = None
    by_fnsku: dict[str, dict] = {}
    by_spec: dict[tuple, dict] = {}
    for r in range(2, (ws.max_row or 1) + 1):
        if ws.cell(r, 2).value not in (None, ""):
            try:
                spp = int(ws.cell(r, 2).value)
            except (TypeError, ValueError):
                pass
        if ws.cell(r, 3).value not in (None, ""):
            size = ws.cell(r, 3).value
        fn = ws.cell(r, 5).value
        if not fn:
            continue
        net = ws.cell(r, 9).value
        gross = ws.cell(r, 10).value
        price_ctn = ws.cell(r, 11).value
        lwh = _parse_lwh(ws.cell(r, 7).value)
        ppc = ws.cell(r, 8).value
        if price_ctn is None:
            price_ctn = _eval_master_price_formula(ws_f.cell(r, 11).value, ppc)
        color = ws.cell(r, 4).value
        rec = {
            "fnsku": str(fn).strip(),
            "strips_per_pack": spp,
            "size": size,
            "size_full": _norm_size(size, short_to_full),
            "color": color,
            "color_canon": _norm_color(color, color_to_canon),
            "packs_per_carton": int(ppc) if ppc else None,
            "net_kg": float(net) if net is not None else None,
            "gross_kg": float(gross) if gross is not None else None,
            "price_per_carton": float(price_ctn) if price_ctn is not None else None,
            "lwh": lwh,
        }
        by_fnsku[str(fn).strip()] = rec
        if rec["strips_per_pack"] and rec["size_full"] and rec["color_canon"]:
            by_spec[(rec["strips_per_pack"], rec["size_full"], rec["color_canon"])] = rec
    wb.close()
    wb_f.close()
    return {
        "by_fnsku": by_fnsku,
        "by_spec": by_spec,
        "short_to_full": short_to_full,
        "color_to_canon": color_to_canon,
    }


def load_master_by_fnsku() -> dict[str, dict]:
    """Compat: FNSKU → 1039/净毛重（不含 C07 同规格继承）。"""
    return load_master_index()["by_fnsku"]


def _spec_cartons(s: dict) -> int:
    v = s.get("qty_cartons")
    if v in (None, ""):
        v = s.get("cartons")
    try:
        return int(v or 0)
    except (TypeError, ValueError):
        return 0


def _lookup_master_row(s: dict, index: dict) -> tuple[dict | None, str]:
    """FNSKU first; else C07 inherit: same 条数/包+尺寸+颜色 → C04 1039 行."""
    fn = str(s.get("fnsku") or "").strip()
    by_fn = index.get("by_fnsku") or {}
    if fn and fn in by_fn:
        return by_fn[fn], "fnsku"
    spp = s.get("strips_per_pack")
    try:
        spp_i = int(spp) if spp not in (None, "") else None
    except (TypeError, ValueError):
        spp_i = None
    size_full = _norm_size(s.get("size"), index.get("short_to_full") or {})
    color = _norm_color(s.get("color"), index.get("color_to_canon") or {})
    if spp_i and size_full and color:
        rec = (index.get("by_spec") or {}).get((spp_i, size_full, color))
        if rec:
            return rec, "c04_same_spec_1039"
    return None, "missing"


def warehouse_money_weights(wh_specs: list[dict], master: dict) -> dict:
    """Aggregate net/gross/declare amount for one warehouse from its specs."""
    index = master if "by_fnsku" in master else {"by_fnsku": master, "by_spec": {}}
    total_net = 0.0
    total_gross = 0.0
    total_amount = 0.0
    missing = []
    joins = []
    lwh_votes: dict[tuple, int] = {}
    for s in wh_specs:
        ctn = _spec_cartons(s)
        m, how = _lookup_master_row(s, index)
        if not m or ctn <= 0:
            missing.append(s)
            continue
        if how != "fnsku":
            joins.append(
                f"{s.get('fnsku')}: {how} via C04 {m.get('fnsku')} ${m.get('price_per_carton')}/ctn"
            )
        total_net += ctn * float(m["net_kg"] or 0)
        total_gross += ctn * float(m["gross_kg"] or 0)
        total_amount += ctn * float(m["price_per_carton"] or 0)
        lwh_votes[m["lwh"]] = lwh_votes.get(m["lwh"], 0) + ctn
    lwh = max(lwh_votes.items(), key=lambda kv: kv[1])[0] if lwh_votes else (57.0, 40.0, 33.0)
    return {
        "net_kg": round(total_net, 2),
        "gross_kg": round(total_gross, 2),
        "declare_amount_raw": round(total_amount, 2),
        "lwh": lwh,
        "missing_specs": missing,
        "spec_joins": joins,
    }


def price_check(amount: float, strips: int) -> tuple[float, float]:
    unit = round(amount / strips, 2) if strips else 0.0
    amount2 = round(unit * strips, 2)
    return unit, amount2


def _reinject_freight_media(target: Path, template: Path) -> None:
    """openpyxl save strips WPS DISPIMG media; copy media/cellimages back from template."""
    with zipfile.ZipFile(target, "r") as zt:
        media = [n for n in zt.namelist() if n.startswith("xl/media/") and not n.endswith("/")]
        if media:
            return
        entries = {name: zt.read(name) for name in zt.namelist()}

    with zipfile.ZipFile(template, "r") as zs:
        template_rels = zs.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        for name in zs.namelist():
            if name.startswith("xl/media/") and not name.endswith("/"):
                entries[name] = zs.read(name)
            elif name in FREIGHT_MEDIA_PARTS:
                entries[name] = zs.read(name)

    ct = entries["[Content_Types].xml"].decode("utf-8")
    if 'Extension="png"' not in ct:
        ct = ct.replace(
            '<Default Extension="rels"',
            '<Default Extension="png" ContentType="image/png"/><Default Extension="rels"',
        )
    if "cellimages.xml" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/xl/cellimages.xml" '
            'ContentType="application/vnd.wps-officedocument.cellimage+xml"/></Types>',
        )
    entries["[Content_Types].xml"] = ct.encode("utf-8")

    rels = entries["xl/_rels/workbook.xml.rels"].decode("utf-8")
    if "cellimages.xml" not in rels:
        m = re.search(
            r'<Relationship[^>]+Target="cellimages\.xml"[^>]*/>',
            template_rels,
        )
        if m:
            rel_line = m.group(0)
            ids = {int(x) for x in re.findall(r'Id="rId(\d+)"', rels)}
            tid = re.search(r'Id="(rId\d+)"', rel_line)
            if tid and tid.group(1) in rels:
                next_id = max(ids) + 1 if ids else 1
                rel_line = re.sub(r'Id="rId\d+"', f'Id="rId{next_id}"', rel_line)
            rels = rels.replace("</Relationships>", rel_line + "</Relationships>")
            entries["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")

    tmp = target.with_suffix(".mediafix.xlsx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zw:
        for name, data in entries.items():
            zw.writestr(name, data)
    tmp.replace(target)


def fill_freight(path: Path, row: dict, money: dict, unit_price: float) -> None:
    """Fill freight 下单模板: only allowed change fields; keep static from template/defaults.

    Change: E3 仓库代码, E7 收件人姓名(空), E8 地址=仓码+地址,
            A15 FBA, B15 追踪码, C15 箱, D15/E15 品名, F15 floor(条/箱), G15 单价,
            S15 样例图片(+media reinject).
    Static: C2 客户单号=111111111; I/J/K/L=16; HS/材质/用途/品牌三列; 勿填 E10–E12.
    """
    del money  # freight 箱重尺不再用总表毛重/箱规
    wb = openpyxl.load_workbook(path)
    ws = wb["下单模板"]
    wh = row["warehouse_code"]
    addr = WH_ADDRESS.get(wh) or WH_ADDRESS.get(str(wh).upper()) or get_address(wh)
    if not addr:
        raise KeyError(f"warehouse address missing: {wh}")
    cartons = int(row["cartons"])
    strips = int(row["strips"])
    per_carton = math.floor(strips / cartons) if cartons else 0

    # 客户单号固定；禁止写入追踪码
    ws["C2"] = FREIGHT_CUSTOMER_NO
    ws["E3"] = wh
    ws["E7"] = None  # 收件人姓名不填
    ws["E8"] = addr["freight_addr"]
    for clear_ref in ("E10", "E11", "E12"):
        ws[clear_ref] = None

    ws["A15"] = row["fba_shipment_id"]
    ws["B15"] = row.get("tracking_id") or ""
    ws["C15"] = cartons
    ws["D15"] = "桌布"
    ws["E15"] = "table cloth"
    ws["F15"] = per_carton
    ws["G15"] = unit_price
    ws["H15"] = "USD"
    dim = FREIGHT_CARTON_DIM_DEFAULT
    ws["I15"] = dim
    ws["J15"] = dim
    ws["K15"] = dim
    ws["L15"] = dim
    ws["M15"] = "6302539090"
    ws["N15"] = "聚酯纤维/polyester"
    ws["O15"] = "餐桌"
    ws["P15"] = "无"
    ws["Q15"] = "无"
    ws["R15"] = "无"
    s15 = ws["S15"].value
    if not (isinstance(s15, str) and "DISPIMG" in s15.upper()):
        ws["S15"] = FREIGHT_DISPIMG
    store = str(row.get("store") or row.get("entity_code") or "C04").upper()
    if store not in FREIGHT_SALES_LINK:
        store = "C04" if store.startswith("C04") else store
    ws["T15"] = FREIGHT_SALES_LINK.get(store, FREIGHT_SALES_LINK["C04"])
    wb.save(path)
    wb.close()
    _reinject_freight_media(path, FREIGHT_TEMPLATE)
    # 强制校验：公式在、media 不在 = 图丢失（FC/AWD 均不可接受）
    import zipfile as _zf

    with _zf.ZipFile(path, "r") as zf:
        media_n = sum(
            1 for n in zf.namelist() if n.startswith("xl/media/") and not n.endswith("/")
        )
    if media_n < 1:
        raise RuntimeError(
            f"freight image missing after reinject: {path} (xl/media={media_n})"
        )


def fill_customs(path: Path, row: dict, money: dict, unit_price: float, amount: float, protocol: str) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb["基础表"]
    cartons = int(row["cartons"])
    strips = int(row["strips"])
    # change fields only
    ws.cell(7, 12).value = money["gross_kg"]  # 总毛重
    ws.cell(8, 6).value = protocol  # 合同协议号
    ws.cell(8, 12).value = money["net_kg"]  # 总净重
    ws.cell(9, 4).value = "许昌波丽丝贸易有限公司"  # 表内抬头
    ws.cell(9, 6).value = "波丽丝"
    ws.cell(9, 12).value = cartons  # 总件数
    # commodity row 14
    ws.cell(14, 3).value = "6302539090"
    ws.cell(14, 4).value = "聚酯纤维桌布"
    ws.cell(14, 6).value = cartons  # 件数
    ws.cell(14, 7).value = strips  # 成交数量
    ws.cell(14, 8).value = "条"
    ws.cell(14, 9).value = unit_price
    ws.cell(14, 10).value = amount
    ws.cell(14, 11).value = money["net_kg"]
    ws.cell(14, 14).value = "许昌波丽丝贸易有限公司"
    ws.cell(14, 15).value = "无品牌"
    ws.cell(14, 16).value = "无货号"
    # buyer
    ws.cell(2, 8).value = "Amazon.com Services, Inc."
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", required=True)
    ap.add_argument("--recognize-json", type=str, default=None)
    ap.add_argument("--headers-json", type=str, default=None)
    ap.add_argument("--order-code", type=str, default=None)
    ap.add_argument("--ship-date", type=str, default=None, help="YYYY-MM-DD")
    ap.add_argument("--leg-id", type=str, default=None)
    ap.add_argument("--out-freight-dir", type=str, default=None)
    ap.add_argument("--out-customs-dir", type=str, default=None)
    args = ap.parse_args()
    project = Path(args.project)
    rec_path = (
        Path(args.recognize_json)
        if args.recognize_json
        else ROOT / "2_Processed_AI" / "drafts" / "C26079-A_fba_recognize.json"
    )
    headers_path = (
        Path(args.headers_json)
        if args.headers_json
        else ROOT / "2_Processed_AI" / "drafts" / "C26079-A_warehouse_headers.json"
    )
    rec = json.loads(rec_path.read_text(encoding="utf-8-sig"))
    headers = json.loads(headers_path.read_text(encoding="utf-8-sig"))
    master = load_master_index()
    WH_ADDRESS.clear()
    WH_ADDRESS.update((load_addresses().get("warehouses") or {}))

    freight_dir = Path(args.out_freight_dir) if args.out_freight_dir else project / "货代入库贴下单表"
    customs_dir = Path(args.out_customs_dir) if args.out_customs_dir else project / "报关资料表"
    freight_dir.mkdir(parents=True, exist_ok=True)
    customs_dir.mkdir(parents=True, exist_ok=True)

    # map warehouse -> specs from recognize
    specs_by_wh = {
        w["warehouse_code"]: w.get("specs") or [] for w in (rec.get("warehouses") or [])
    }

    order_code = str(args.order_code or headers.get("order_code") or "EF-UNKNOWN")
    leg_id = str(args.leg_id or headers.get("split_leg_id") or f"{order_code}-A")
    ship_date = str(args.ship_date or headers.get("ship_date") or "2026-07-31")
    try:
        y, m, d = ship_date[:10].split("-")
        ship_disp = f"{y}.{int(m)}.{int(d)}"
    except Exception:
        ship_disp = ship_date

    summary = {
        "generated_at": _now(),
        "order_code": order_code,
        "split_leg_id": leg_id,
        "ship_date": ship_date,
        "pricing_method": "per_warehouse: raw=sum(cartons*master_1039_ctn_price); unit=round(raw/strips,2); amount=round(unit*strips,2)",
        "note": "草稿；未进 Final",
        "remaining_for_B": (rec.get("plan_compare") or {}).get("missing_specs_in_fba"),
        "files": [],
        "warnings": [],
    }

    if not FREIGHT_TEMPLATE.is_file() or not CUSTOMS_TEMPLATE.is_file():
        print("missing templates", file=sys.stderr)
        return 1

    grand_amount = 0.0
    grand_raw = 0.0
    store = "C07" if "-C07-" in order_code or order_code.startswith("C07") else "C04"
    for row in headers.get("rows") or []:
        row = {**row, "store": store}
        wh = row["warehouse_code"]
        addr = get_address(wh) or WH_ADDRESS.get(str(wh).upper())
        if not addr:
            addr = wh_placeholder(wh)
            WH_ADDRESS[str(wh).upper()] = addr
            summary["warnings"].append(
                f"{wh}: address not in amazon_warehouse_addresses.json; placeholder used"
            )
        else:
            WH_ADDRESS[str(wh).upper()] = addr
        money = warehouse_money_weights(specs_by_wh.get(wh, []), master)
        if money.get("spec_joins"):
            summary["warnings"].extend(money["spec_joins"])
        if money["missing_specs"]:
            summary["warnings"].append(f"{wh}: {len(money['missing_specs'])} specs missing master join")
        raw = float(money["declare_amount_raw"])
        unit, amount = price_check(raw, int(row["strips"]))
        if amount <= 0:
            summary["warnings"].append(f"{wh}: declare amount 0 — check master join / specs")
        grand_amount += amount
        grand_raw += raw

        freight_name = f"货代下单表_{leg_id} {ship_disp}桌布 {wh} {row['cartons']}箱.xlsx"
        customs_name = (
            f"畅选+波丽丝 {ship_disp}发货 {leg_id} {wh} {row['cartons']}箱 市场采购数据.xlsx"
        )
        freight_path = freight_dir / freight_name
        customs_path = customs_dir / customs_name
        shutil.copy2(FREIGHT_TEMPLATE, freight_path)
        shutil.copy2(CUSTOMS_TEMPLATE, customs_path)

        fill_freight(freight_path, row, money, unit)
        protocol = allocate_protocol_no(
            ship_date[:10],
            project=project.name,
            warehouse_code=wh,
            order_or_plan=str(leg_id or order_code),
        )
        fill_customs(customs_path, row, money, unit, amount, protocol)

        summary["files"].append(
            {
                "warehouse_code": wh,
                "fba_shipment_id": row["fba_shipment_id"],
                "tracking_id": row.get("tracking_id"),
                "cartons": row["cartons"],
                "strips": row["strips"],
                "raw_carton_x_price": round(raw, 2),
                "unit_price_usd": unit,
                "declare_amount_usd": amount,
                "net_kg": money["net_kg"],
                "gross_kg": money["gross_kg"],
                "protocol_no": protocol,
                "freight_file": str(freight_path),
                "customs_file": str(customs_path),
            }
        )
        print(
            f"{wh}: cartons={row['cartons']} strips={row['strips']} "
            f"raw={round(raw,2)} unit={unit} amount={amount}"
        )

    summary["total_raw_carton_x_price"] = round(grand_raw, 2)
    summary["total_declare_amount_usd"] = round(grand_amount, 2)
    # 规格汇总 A vs 分仓 B：误差率闸门（≤1% 不报告）
    spec_amt = round(grand_raw, 2)  # raw 合计 = 规格汇总货值（同算法）
    # 更精确：从全部 specs 重算 A（与 grand_raw 应一致）
    a_amt = a_net = a_gross = 0.0
    for specs in specs_by_wh.values():
        mny = warehouse_money_weights(specs, master)
        a_amt += float(mny["declare_amount_raw"])
        a_net += float(mny["net_kg"])
        a_gross += float(mny["gross_kg"])
    b_net = sum(float(f["net_kg"]) for f in summary["files"])
    b_gross = sum(float(f["gross_kg"]) for f in summary["files"])

    def _err_pct(b: float, a: float) -> float | None:
        if a == 0:
            return 0.0 if b == 0 else None
        return abs(b - a) / a * 100.0

    er_amt = _err_pct(grand_amount, a_amt)
    er_net = _err_pct(b_net, a_net)
    er_gross = _err_pct(b_gross, a_gross)
    summary["reconcile"] = {
        "doc": "Templates/Amount_Weight_Reconcile.md",
        "A_spec": {"amount": round(a_amt, 2), "net_kg": round(a_net, 2), "gross_kg": round(a_gross, 2)},
        "B_warehouse": {
            "amount": round(grand_amount, 2),
            "net_kg": round(b_net, 2),
            "gross_kg": round(b_gross, 2),
        },
        "error_rate_pct": {
            "declare": None if er_amt is None else round(er_amt, 4),
            "net": None if er_net is None else round(er_net, 4),
            "gross": None if er_gross is None else round(er_gross, 4),
        },
        "pass_max_pct": 1.0,
        "gate": {
            "declare": "PASS" if er_amt is not None and er_amt <= 1.0 else "CONFIRM",
            "net": "PASS" if er_net is not None and er_net <= 1.0 else "CONFIRM",
            "gross": "PASS" if er_gross is not None and er_gross <= 1.0 else "CONFIRM",
        },
    }
    summary["reconcile"]["gate"]["overall"] = (
        "PASS"
        if all(summary["reconcile"]["gate"][k] == "PASS" for k in ("declare", "net", "gross"))
        else "CONFIRM"
    )
    print(
        f"五仓合计 raw箱×单价={summary['total_raw_carton_x_price']} "
        f"正算反算货值={summary['total_declare_amount_usd']}"
    )
    print(
        f"对照误差率 declare={summary['reconcile']['error_rate_pct']['declare']}% "
        f"net={summary['reconcile']['error_rate_pct']['net']}% "
        f"gross={summary['reconcile']['error_rate_pct']['gross']}% "
        f"gate={summary['reconcile']['gate']['overall']}"
    )
    if summary["reconcile"]["gate"]["overall"] != "PASS":
        summary["warnings"].append(
            "误差率>1%：须人工确认后再继续（见 Templates/Amount_Weight_Reconcile.md）"
        )
        print("WARN: 误差率超过 1%，须人工确认")

    out = ROOT / "2_Processed_AI" / "drafts" / f"{leg_id}_freight_customs_summary.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {out}")
    for w in summary["warnings"]:
        print("WARN:", w)

    # Optional leftover note only when recognize payload carries remaining_for_B
    remaining = summary.get("remaining_for_B")
    if remaining:
        note_dir = project / "拆单比对"
        note_dir.mkdir(parents=True, exist_ok=True)
        note_path = note_dir / f"{leg_id}_REMAINING.md"
        note_path.write_text(
            json.dumps({"remaining_for_B": remaining}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(f"wrote {note_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
